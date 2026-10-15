from __future__ import annotations

from io import BytesIO
from datetime import date

from openpyxl import Workbook, load_workbook

from app.core import db
from app.exporters.requirements_xlsx import render_requirement_workbook
from app.models import Attachment, Bid, Evaluation, Requirement
from app.parsers.excel import read_sheets
from app.parsers.requirements import normalize


def test_excel_reader_finds_header_on_row_three_and_aliases():
    workbook = Workbook()
    cover = workbook.active
    assert cover is not None
    cover.title = "表紙"
    cover.append(["架空の要領"])
    sheet = workbook.create_sheet("別紙1")
    sheet.append(["説明"])
    sheet.append([])
    sheet.append(["No", "要件名", "要件内容", "優先度", "備考"])
    sheet.append(["001", "ログイン", "利用者がログインできる", "必須", "顧客記入"])
    output = BytesIO()
    workbook.save(output)

    sheets = read_sheets(output.getvalue(), "requirements.xlsx")
    rows = normalize(sheets)

    assert len(rows) == 1
    assert rows[0]["internal_code"] == "001"
    assert rows[0]["content"] == "利用者がログインできる"
    assert rows[0]["priority"] == "必須"
    assert rows[0]["source_sheet"] == "別紙1"
    assert rows[0]["source_row"] == 4


def test_tokai_real_requirements_are_41_with_customer_codes():
    """Đợt 14 (B1/B2): tệp Tokai thật → đúng 41 requirement, giữ mã khách.

    Sheet bìa 表紙 (không có cột mã) bị bỏ; cột 要件ID thắng cột No. khi
    tìm cột mã; mã khách làm internal_code để truy vết và xuất bảng đáp ứng.
    """
    from pathlib import Path

    real_dir = Path(__file__).resolve().parents[2] / "sample_data" / "rfp_tokai_minato" / "real"
    xlsx = next(real_dir.glob("*.xlsx"))
    sheets = read_sheets(xlsx.read_bytes(), xlsx.name)
    rows = normalize(sheets)

    assert len(rows) == 41
    assert {r["source_sheet"] for r in rows} == {"機能要件", "非機能要件"}
    assert [r["internal_code"] for r in rows[:3]] == ["FR-QA-001", "FR-QA-002", "FR-QA-003"]
    functional = [r for r in rows if r["source_sheet"] == "機能要件"]
    non_functional = [r for r in rows if r["source_sheet"] == "非機能要件"]
    assert len(functional) == 25
    assert len(non_functional) == 16
    assert all(r["internal_code"].startswith("FR-") for r in functional)
    assert all(r["internal_code"].startswith("NFR-") for r in non_functional)
    assert len({r["internal_code"] for r in rows}) == 41


def test_normalize_skips_cover_sheet_without_code_column():
    """Sheet bìa có nội dung nhưng không có cột mã → bị bỏ khi file có sheet
    khác có cột mã; nếu cả file không có cột mã thì vẫn parse (REQ-xxx)."""
    sheets = {
        "表紙": [{"column_1": "記入方法", "column_2": "黄色の列に記入"}],
        "要件": [{"要件ID": "FR-001", "要件内容": "ログインできる"}],
    }
    rows = normalize(sheets)
    assert len(rows) == 1
    assert rows[0]["internal_code"] == "FR-001"
    assert rows[0]["source_sheet"] == "要件"

    # Cả file không có cột mã → fallback REQ-xxx, không bỏ sheet.
    rows = normalize({"A": [{"内容": "Yêu cầu A"}], "B": [{"内容": "Yêu cầu B"}]})
    assert len(rows) == 2
    assert all(r["internal_code"].startswith("REQ-") for r in rows)


def test_code_column_prefers_exact_match_over_no_column():
    """Cột 'No.' (số thứ tự) không được thắng cột '要件ID' khi tìm cột mã."""
    sheets = {
        "要件": [{"No.": "1", "要件ID": "FR-QA-001", "要件内容": "質問応答"}],
    }
    rows = normalize(sheets)
    assert rows[0]["internal_code"] == "FR-QA-001"


def test_requirement_export_writes_response_and_note_without_overwriting_remark(tmp_path):
    source = tmp_path / "customer.xlsx"
    workbook = Workbook()
    active = workbook.active
    assert active is not None
    workbook.remove(active)
    sheet = workbook.create_sheet("別紙1")
    sheet.append(["表紙"])
    sheet.append([])
    sheet.append(["No", "要件内容", "備考"])
    sheet.append(["001", "CRM", "顧客の備考"])
    workbook.save(source)

    session = db.SessionLocal()
    bid = Bid(project_name="架空", customer_name="架空", submission_deadline=date(2026, 12, 1), customer_language="ja", request_content="架空", created_by="test")
    session.add(bid)
    session.commit()
    session.add(Attachment(bid_id=bid.id, filename="customer.xlsx", file_type="requirement", detected_type="requirement", size_bytes=source.stat().st_size, path=str(source), original_path=str(source), parse_status="ok"))
    requirement = Requirement(bid_id=bid.id, internal_code="001", content="CRM", source_sheet="別紙1", source_row=4, source_cols={"__header_row__": 3})
    session.add(requirement)
    session.commit()
    session.add(Evaluation(bid_id=bid.id, requirement_id=requirement.id, rating="△", needs_clarification=False, note_vi="要件確認が必要", batch_no=1, solution_direction="rag"))
    session.commit()

    _, data = render_requirement_workbook(session, bid.id, "ja")
    result = load_workbook(BytesIO(data))
    output_sheet = result["別紙1"]
    headers = [output_sheet.cell(3, column).value for column in range(1, output_sheet.max_column + 1)]
    assert "対応可否" in headers
    assert "補足説明" in headers
    assert output_sheet.cell(4, headers.index("対応可否") + 1).value == "△"
    assert output_sheet.cell(4, headers.index("補足説明") + 1).value == "要件確認が必要"
    assert output_sheet.cell(4, 3).value == "顧客の備考"
    session.close()


def test_requirement_export_uses_requested_note_language(tmp_path):
    source = tmp_path / "customer.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Requirements"
    sheet.append(["No", "Requirement"])
    sheet.append(["001", "CRM"])
    workbook.save(source)

    session = db.SessionLocal()
    bid = Bid(
        project_name="架空",
        customer_name="架空",
        submission_deadline=date(2026, 12, 1),
        customer_language="ja",
        request_content="架空の要件本文です。",
        created_by="test",
    )
    session.add(bid)
    session.commit()
    session.add(
        Attachment(
            bid_id=bid.id,
            filename="customer.xlsx",
            file_type="requirement",
            detected_type="requirement",
            size_bytes=source.stat().st_size,
            path=str(source),
            original_path=str(source),
            parse_status="ok",
        )
    )
    requirement = Requirement(
        bid_id=bid.id,
        internal_code="001",
        content="CRM",
        source_sheet="Requirements",
        source_row=2,
        source_cols={"__header_row__": 1},
    )
    session.add(requirement)
    session.commit()

    _, data = render_requirement_workbook(session, bid.id, "ja", note_language="en")
    result = load_workbook(BytesIO(data))
    output_sheet = result["Requirements"]
    headers = [output_sheet.cell(1, column).value for column in range(1, output_sheet.max_column + 1)]
    assert output_sheet.cell(2, headers.index("補足説明") + 1).value == "Coverage: chưa có kết quả đánh giá"
    session.close()
