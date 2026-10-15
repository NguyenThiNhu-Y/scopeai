from io import BytesIO
from types import SimpleNamespace

import openpyxl
from docx import Document
from pptx import Presentation

from app.exporters.estimate_xlsx import render_estimate
from app.exporters.proposal_docx import render_proposal
from app.exporters.proposal_pptx import render_slides


def _unapproved():
    return {
        "project_name": "Dự án giả lập",
        "customer_name": "Khách hàng giả lập",
        "currency": "VND",
        "quotation": "[Presales điền]",
        "approved": False,
        "total": None,
        "quotation_detail": {"contract_model": "fixed_price", "by_role": {"backend": {"man_day": 3.0}}, "phase_man_days": {"poc": 3.0}},
        "sections": {"conclusion": "Kết luận [Presales điền]"},
        "wbs": [SimpleNamespace(task_code="1.1", name_vi="Việc", phase="poc", parent_code=None, role="backend", man_day=3)],
        "schedule": None,
        "questions": [],
    }


def _all_cell_text(workbook) -> str:
    return " ".join(str(cell.value) for sheet in workbook for row in sheet.iter_rows() for cell in row if cell.value is not None)


def test_unapproved_excel_uses_placeholder_not_real_amount():
    workbook = openpyxl.load_workbook(BytesIO(render_estimate(_unapproved())), data_only=False)
    assert workbook["Tóm tắt"]["B10"].value == "[Presales điền]"
    sheet = workbook["Báo giá"]
    amounts = [sheet.cell(r, 4).value for r in range(1, sheet.max_row + 1) if sheet.cell(r, 4).value is not None]
    assert amounts and all(value in ("[Presales điền]", "Thành tiền") for value in amounts)
    assert "123456" not in _all_cell_text(workbook)


def test_approved_flag_false_hides_detail_even_if_amounts_leak_into_data():
    data = _unapproved()
    data["quotation_detail"] = {"contract_model": "fixed_price", "by_role": {"backend": {"man_day": 3.0, "rate": 41152, "amount": 123456}}, "grand_total": 123456, "by_phase": {"poc": 123456}}
    workbook = openpyxl.load_workbook(BytesIO(render_estimate(data)), data_only=False)
    assert "123456" not in _all_cell_text(workbook)
    assert "41152" not in _all_cell_text(workbook)
    document = Document(BytesIO(render_proposal(data)))
    assert "123,456" not in "\n".join(p.text for p in document.paragraphs)
    deck = Presentation(BytesIO(render_slides(data)))
    deck_text = " ".join(shape.text_frame.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
    assert "123,456" not in deck_text and "[Presales điền]" in deck_text
