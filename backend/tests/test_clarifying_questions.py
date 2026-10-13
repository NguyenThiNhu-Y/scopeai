"""Tests cho luồng hỏi làm rõ — FR-010/FR-011 (BR-02, BR-04, BR-07, NFR-004)."""

from __future__ import annotations

from app.core import db
from app.llm.mock_client import MockLLMClient
from app.models import PipelineStep
from openpyxl import load_workbook
from io import BytesIO
from tests.helpers import bid_status, create_bid, step_status, wait_until

_JAPANESE_RANGES = ((0x3040, 0x309F), (0x30A0, 0x30FF))
_VIETNAMESE_CHARS = (
    "àáảãạăắằẳẵặâấầẩẫậèéẻẽẹêếềểễệìíỉĩịòóỏõọôốồổỗộơớờởỡợùúủũụưứừửữựỳýỷỹỵđ"
)


def _has_japanese(text: str) -> bool:
    return any(
        start <= ord(ch) <= end for start, end in _JAPANESE_RANGES for ch in text
    )


def _has_vietnamese(text: str) -> bool:
    return any(ch in _VIETNAMESE_CHARS for ch in text)


def _start_pipeline(client, bid_id: int):
    resp = client.post(f"/api/bids/{bid_id}/pipeline/start")
    assert resp.status_code == 202, resp.text
    return resp.json()


def _get_questions(client, bid_id: int, round_no: int | None = None) -> dict:
    params = {"round": round_no} if round_no else None
    resp = client.get(f"/api/bids/{bid_id}/questions", params=params)
    assert resp.status_code == 200, resp.text
    return resp.json()


def _answer_round(client, bid_id: int, round_no: int, assumption_ids: set[int] | None = None):
    assumption_ids = assumption_ids or set()
    data = _get_questions(client, bid_id, round_no)
    answers = []
    for q in data["items"]:
        answers.append(
            {
                "question_id": q["id"],
                "answer": "約100万件（仮の回答）" if round_no == 1 else "Giả định: dùng RAG với chunk 512",
                "is_assumption": q["id"] in assumption_ids,
            }
        )
    return client.post(
        f"/api/bids/{bid_id}/questions/{round_no}/answers",
        json={"answers": answers},
        headers={"X-User": "Nhu"},
    )


def test_gap_questions_in_customer_language(client, mock_llm):
    """BR-07/NFR-004: khách ja → câu hỏi tiếng Nhật + reason_vi tiếng Việt, ≤ 7 câu."""
    bid_id = create_bid(client, customer_language="ja")
    _start_pipeline(client, bid_id)

    assert wait_until(lambda: bid_status(client, bid_id) == "waiting_clarification")

    data = _get_questions(client, bid_id)
    items = data["items"]
    assert 1 <= len(items) <= 7
    assert data["rounds_used"] == 1
    assert data["max_rounds"] == 2

    for q in items:
        assert _has_japanese(q["question"]), f"Câu hỏi không phải tiếng Nhật: {q['question']!r}"
        assert _has_vietnamese(q["reason_vi"]), f"Lý do không phải tiếng Việt: {q['reason_vi']!r}"
        assert q["impact"] in ("solution", "effort", "price")
        assert q["related_req_codes"] == []
        assert q["status"] == "awaiting_answer"
        assert q["answer"] is None
        assert q["is_assumption"] is False

    # Trạng thái bước gap = waiting_clarification (FR-010.3).
    assert step_status(client, bid_id, "gap") == "waiting_clarification"


def test_questions_export_is_customer_language_and_answer_column_is_yellow(client, mock_llm):
    bid_id = create_bid(client, customer_language="ja")
    _start_pipeline(client, bid_id)
    assert wait_until(lambda: bid_status(client, bid_id) == "waiting_clarification")

    response = client.get(f"/api/bids/{bid_id}/questions/export")
    assert response.status_code == 200
    sheet = load_workbook(BytesIO(response.content)).active
    assert sheet.title == "Q&A"
    assert [sheet.cell(1, column).value for column in range(1, 6)] == ["質問コード", "ラウンド", "質問", "理由", "回答"]
    assert any("ご要望" in str(sheet.cell(row, 3).value) for row in range(2, sheet.max_row + 1))
    assert sheet.cell(2, 5).fill.fgColor.rgb in {"FFFFFF00", "00FFFF00"}


def test_questions_import_merges_answers_and_resumes_pipeline(client, mock_llm):
    bid_id = create_bid(client, customer_language="ja")
    _start_pipeline(client, bid_id)
    assert wait_until(lambda: bid_status(client, bid_id) == "waiting_clarification")
    exported = client.get(f"/api/bids/{bid_id}/questions/export")
    workbook = load_workbook(BytesIO(exported.content))
    sheet = workbook.active
    for row in range(2, sheet.max_row + 1):
        sheet.cell(row, 5).value = "回答済み"
    output = BytesIO()
    workbook.save(output)
    response = client.post(
        f"/api/bids/{bid_id}/questions/import",
        files={"file": ("answers.xlsx", output.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 200, response.text
    assert response.json()["pipeline_resumed_from"] == "gap"
    assert wait_until(lambda: bid_status(client, bid_id) in {"analyzing", "waiting_clarification", "waiting_tech_approval"})


def test_questions_import_rejects_missing_round_column(client, mock_llm):
    bid_id = create_bid(client, customer_language="ja")
    _start_pipeline(client, bid_id)
    assert wait_until(lambda: bid_status(client, bid_id) == "waiting_clarification")
    workbook = load_workbook(BytesIO(client.get(f"/api/bids/{bid_id}/questions/export").content))
    sheet = workbook.active
    sheet.delete_cols(2)
    output = BytesIO()
    workbook.save(output)

    response = client.post(
        f"/api/bids/{bid_id}/questions/import",
        files={"file": ("answers.xlsx", output.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )

    assert response.status_code == 400
    assert "thiếu cột bắt buộc" in response.json()["error"]["message"]


def test_questions_import_skips_old_and_unknown_codes_without_duplicates(client, mock_llm):
    bid_id = create_bid(client, customer_language="ja")
    _start_pipeline(client, bid_id)
    assert wait_until(lambda: bid_status(client, bid_id) == "waiting_clarification")
    exported = client.get(f"/api/bids/{bid_id}/questions/export")
    workbook = load_workbook(BytesIO(exported.content))
    sheet = workbook.active
    sheet.cell(2, 5).value = "回答済み"
    sheet.cell(3, 1).value = 999999
    sheet.cell(3, 5).value = "存在しない質問"
    output = BytesIO()
    workbook.save(output)

    first = client.post(
        f"/api/bids/{bid_id}/questions/import",
        files={"file": ("answers.xlsx", output.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert first.status_code == 200, first.text
    assert first.json()["warnings"]

    second = client.post(
        f"/api/bids/{bid_id}/questions/import",
        files={"file": ("answers.xlsx", output.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert second.status_code == 200, second.text
    assert "đã đóng" in " ".join(second.json()["warnings"])


def test_questions_import_rejects_more_than_200_rows(client, mock_llm):
    bid_id = create_bid(client, customer_language="ja")
    workbook = load_workbook(BytesIO(client.get(f"/api/bids/{bid_id}/questions/export").content))
    sheet = workbook.active
    for row in range(2, 204):
        sheet.cell(row, 1).value = 1
        sheet.cell(row, 2).value = 1
        sheet.cell(row, 5).value = "回答"
    output = BytesIO()
    workbook.save(output)

    response = client.post(
        f"/api/bids/{bid_id}/questions/import",
        files={"file": ("answers.xlsx", output.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )

    assert response.status_code == 400
    assert "200 dòng" in response.json()["error"]["message"]


def test_questions_import_rejects_more_than_5_mb(client):
    bid_id = create_bid(client, customer_language="ja")
    response = client.post(
        f"/api/bids/{bid_id}/questions/import",
        files={"file": ("answers.xlsx", b"x" * (5 * 1024 * 1024 + 1), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )

    assert response.status_code == 400
    assert "5 MB" in response.json()["error"]["message"]


def test_questions_import_supports_second_round_with_round_column(client, mock_llm):
    bid_id = create_bid(client, customer_language="ja")
    _start_pipeline(client, bid_id)
    assert wait_until(lambda: bid_status(client, bid_id) == "waiting_clarification")
    assert _get_questions(client, bid_id, 1)["items"]

    exported = client.get(f"/api/bids/{bid_id}/questions/export")
    workbook = load_workbook(BytesIO(exported.content))
    sheet = workbook.active
    for row in range(2, sheet.max_row + 1):
        sheet.cell(row, 5).value = "回答済み"
    output = BytesIO()
    workbook.save(output)
    response = client.post(
        f"/api/bids/{bid_id}/questions/import",
        files={"file": ("answers.xlsx", output.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 200, response.text
    assert wait_until(lambda: bid_status(client, bid_id) in {"waiting_clarification", "waiting_tech_approval"})

    second_round = _get_questions(client, bid_id, 2)["items"]
    if not second_round:
        return
    exported = client.get(f"/api/bids/{bid_id}/questions/export")
    workbook = load_workbook(BytesIO(exported.content))
    sheet = workbook.active
    for row in range(2, sheet.max_row + 1):
        if sheet.cell(row, 2).value == 2:
            sheet.cell(row, 5).value = "回答 vòng hai"
    output = BytesIO()
    workbook.save(output)
    response = client.post(
        f"/api/bids/{bid_id}/questions/import",
        files={"file": ("answers.xlsx", output.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 200, response.text


def test_customer_email_draft_is_structured_customer_language_and_editable(client, mock_llm):
    bid_id = create_bid(client, customer_language="ja")
    response = client.post(f"/api/bids/{bid_id}/questions/email-draft", headers={"X-User": "Nhu"})
    assert response.status_code == 200, response.text
    data = response.json()
    assert data["customer_name"] == "Công ty XYZ (giả lập)"
    assert data["sender"] == "Nhu"
    assert data["customer_language"] == "ja"
    assert data["subject"]
    assert data["body"]
    assert "customer_name" not in data["body"]


def test_no_questions_when_not_needed(client, mock_llm):
    """needs_clarification=false → questions rỗng → pipeline chạy tiếp → done."""
    mock_llm["client"] = MockLLMClient(fixture_modes={"GapQuestions": "no_questions"})
    bid_id = create_bid(client)
    _start_pipeline(client, bid_id)

    assert wait_until(lambda: bid_status(client, bid_id) == "waiting_tech_approval")

    data = _get_questions(client, bid_id)
    assert data["items"] == []
    assert data["rounds_used"] == 0
    assert step_status(client, bid_id, "gap") == "completed"
    assert step_status(client, bid_id, "intake") == "completed"


def test_retry_bad_schema_then_fail(client, mock_llm):
    """BR-02: schema xấu → retry ≤ 2 (3 lần gọi) → bước gap failed."""
    mock_llm["client"] = MockLLMClient(fixture_modes={"GapQuestions": "bad_schema"})
    bid_id = create_bid(client)
    _start_pipeline(client, bid_id)

    assert wait_until(lambda: step_status(client, bid_id, "gap") == "failed")

    client_llm = mock_llm["client"]
    assert client_llm.call_counts["GapQuestions"] == 3  # 1 + 2 retry

    session = db.SessionLocal()
    try:
        step = (
            session.query(PipelineStep)
            .filter(PipelineStep.bid_id == bid_id, PipelineStep.step_key == "gap")
            .one()
        )
        assert step.error_message and "GapQuestions" in step.error_message
    finally:
        session.close()


def test_too_many_questions_truncated(client, mock_llm):
    """Vượt 7 câu → giữ 7 câu đầu + cảnh báo (không fail)."""
    mock_llm["client"] = MockLLMClient(fixture_modes={"GapQuestions": "too_many"})
    bid_id = create_bid(client)
    _start_pipeline(client, bid_id)

    assert wait_until(lambda: bid_status(client, bid_id) == "waiting_clarification")
    data = _get_questions(client, bid_id)
    assert len(data["items"]) == 7

    session = db.SessionLocal()
    try:
        step = (
            session.query(PipelineStep)
            .filter(PipelineStep.bid_id == bid_id, PipelineStep.step_key == "gap")
            .one()
        )
        assert step.output_json and "warning" in step.output_json
    finally:
        session.close()


def test_submit_missing_answers_returns_missing_ids(client, mock_llm):
    """FR-011.3: trả lời thiếu câu → 400 kèm details.missing_ids."""
    bid_id = create_bid(client)
    _start_pipeline(client, bid_id)
    assert wait_until(lambda: bid_status(client, bid_id) == "waiting_clarification")

    data = _get_questions(client, bid_id, round_no=1)
    all_ids = {q["id"] for q in data["items"]}
    answered_id = next(iter(all_ids))

    resp = client.post(
        f"/api/bids/{bid_id}/questions/1/answers",
        json={
            "answers": [
                {"question_id": answered_id, "answer": "trả lời một phần", "is_assumption": False}
            ]
        },
    )
    assert resp.status_code == 400
    body = resp.json()["error"]
    assert body["code"] == "VALIDATION_ERROR"
    missing = set(body["details"]["missing_ids"])
    assert missing == all_ids - {answered_id}


def test_submit_answers_on_draft_bid_rejected(client, mock_llm):
    """H-1 (review): answers trên bid KHÔNG ở waiting_clarification → 400, KHÔNG đổi trạng thái."""
    bid_id = create_bid(client)
    assert bid_status(client, bid_id) == "draft"

    resp = client.post(
        f"/api/bids/{bid_id}/questions/1/answers",
        json={"answers": []},
    )
    assert resp.status_code == 400
    body = resp.json()["error"]
    assert body["code"] == "VALIDATION_ERROR"
    assert "chờ làm rõ" in body["message"]

    # Trạng thái hồ sơ KHÔNG bị hỏng (trước fix: bị đổi thành analyzing vĩnh viễn).
    assert bid_status(client, bid_id) == "draft"


def test_submit_answers_empty_on_waiting_bid_returns_missing_ids(client, mock_llm):
    """H-1 (review): answers rỗng khi đang chờ làm rõ → 400 missing_ids (không lọt qua)."""
    bid_id = create_bid(client)
    _start_pipeline(client, bid_id)
    assert wait_until(lambda: bid_status(client, bid_id) == "waiting_clarification")

    resp = client.post(
        f"/api/bids/{bid_id}/questions/1/answers",
        json={"answers": []},
    )
    assert resp.status_code == 400
    body = resp.json()["error"]
    assert body["code"] == "VALIDATION_ERROR"
    assert body["details"]["missing_ids"]
    assert bid_status(client, bid_id) == "waiting_clarification"


def test_submit_answers_wrong_round_rejected(client, mock_llm):
    """H-1 (review): trả lời vòng không có câu hỏi chờ → 400 rõ ràng."""
    bid_id = create_bid(client)
    _start_pipeline(client, bid_id)
    assert wait_until(lambda: bid_status(client, bid_id) == "waiting_clarification")

    resp = client.post(
        f"/api/bids/{bid_id}/questions/2/answers",
        json={"answers": []},
    )
    assert resp.status_code == 400
    body = resp.json()["error"]
    assert body["code"] == "VALIDATION_ERROR"
    assert "Không có câu hỏi chờ trả lời" in body["message"]
    assert bid_status(client, bid_id) == "waiting_clarification"


def test_assumption_answer_marked(client, mock_llm):
    """FR-011.2: trả lời giả định → is_assumption=true + status answered."""
    bid_id = create_bid(client)
    _start_pipeline(client, bid_id)
    assert wait_until(lambda: bid_status(client, bid_id) == "waiting_clarification")

    data = _get_questions(client, bid_id, round_no=1)
    assumption_id = data["items"][0]["id"]

    resp = _answer_round(client, bid_id, 1, assumption_ids={assumption_id})
    assert resp.status_code == 200
    assert resp.json()["status"] == "analyzing"
    assert resp.json()["pipeline_resumed_from"] == "gap"

    data = _get_questions(client, bid_id, round_no=1)
    by_id = {q["id"]: q for q in data["items"]}
    assert by_id[assumption_id]["is_assumption"] is True
    assert by_id[assumption_id]["status"] == "answered"
    assert by_id[assumption_id]["answer"]
    other = next(q for q in data["items"] if q["id"] != assumption_id)
    assert other["is_assumption"] is False
    assert other["status"] == "answered"


def test_round_two_no_llm_call(client, mock_llm):
    """FR-010.4: sau vòng 2, code KHÔNG gọi LLM nữa (call_count không tăng).

    Đợt 10: mock chỉ hỏi vòng 2 khi còn thông tin chặn — dùng câu trả lời
    chứa từ khóa chặn (未定) để kích hoạt vòng 2.
    """
    bid_id = create_bid(client)
    _start_pipeline(client, bid_id)
    assert wait_until(lambda: bid_status(client, bid_id) == "waiting_clarification")

    client_llm = mock_llm["client"]
    assert client_llm.call_counts["GapQuestions"] == 1

    # Vòng 1: trả lời còn thông tin chặn → sinh câu hỏi vòng 2.
    data = _get_questions(client, bid_id, 1)
    answers = [
        {"question_id": q["id"], "answer": "未定（後日回答）", "is_assumption": False}
        for q in data["items"]
    ]
    resp = client.post(
        f"/api/bids/{bid_id}/questions/1/answers",
        json={"answers": answers},
        headers={"X-User": "Nhu"},
    )
    assert resp.status_code == 200
    assert wait_until(lambda: bid_status(client, bid_id) == "waiting_clarification")
    assert client_llm.call_counts["GapQuestions"] == 2

    data = _get_questions(client, bid_id)
    assert data["rounds_used"] == 2
    round2_ids = {q["id"] for q in data["items"] if q["round_no"] == 2}
    assert round2_ids

    # Vòng 2: trả lời đủ → resume từ gap → rounds_used=2 → KHÔNG gọi LLM → done.
    resp = _answer_round(client, bid_id, 2)
    assert resp.status_code == 200
    assert wait_until(lambda: bid_status(client, bid_id) == "waiting_tech_approval")

    assert client_llm.call_counts["GapQuestions"] == 2  # không tăng
    assert client_llm.call_counts["IntakeExtraction"] == 1  # intake không chạy lại


def test_round_two_response_exposes_blocking_reason(client, mock_llm):
    """UI có lý do rõ ràng khi vòng 2 thực sự còn thông tin chặn."""
    bid_id = create_bid(client)
    _start_pipeline(client, bid_id)
    assert wait_until(lambda: bid_status(client, bid_id) == "waiting_clarification")

    data = _get_questions(client, bid_id, 1)
    answers = [
        {"question_id": q["id"], "answer": "未定（後日回答）", "is_assumption": False}
        for q in data["items"]
    ]
    response = client.post(
        f"/api/bids/{bid_id}/questions/1/answers",
        json={"answers": answers},
        headers={"X-User": "Nhu"},
    )
    assert response.status_code == 200
    assert wait_until(lambda: bid_status(client, bid_id) == "waiting_clarification")

    data = _get_questions(client, bid_id)
    assert data["rounds_used"] == 2
    assert data["round_two_reason_vi"]
