from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from app.llm.mock_client import MockLLMClient
from app.schemas.llm_schemas.evaluation import EvaluationBatch
from app.schemas.llm_schemas.intake_extraction import IntakeExtraction
from app.schemas.llm_schemas.solution_proposal import SolutionProposal
from app.schemas.llm_schemas.proposal import ProposalOutput


def test_japanese_sample_has_36_distinct_fr_nfr_requirements():
    workbook = load_workbook(Path(__file__).resolve().parents[2] / "sample_data" / "ja_requirements.xlsx", read_only=True)
    rows = list(workbook["別紙1"].iter_rows(min_row=4, values_only=True))
    codes = [row[0] for row in rows]
    contents = [row[3] for row in rows]
    assert len(rows) == 36
    assert all(str(code).startswith(("FR-", "NFR-")) for code in codes)
    assert len(set(codes)) == len(codes)
    assert len(set(contents)) == len(contents)


def test_mock_evaluation_varies_by_requirement_content():
    client = MockLLMClient()
    result = client.complete(
        schema=EvaluationBatch,
        prompt="evaluate",
        context={
            "batch_codes": ["FR-001", "NFR-001", "FR-002", "NFR-002", "FR-003", "NFR-003"],
            "requirements": [
                {"req_code": "FR-001", "content": "顧客マスタを登録・更新できる", "category": "顧客管理"},
                {"req_code": "NFR-001", "content": "監査ログを保持しアクセス制御する", "category": "セキュリティ"},
                {"req_code": "FR-002", "content": "全文検索と条件検索を行う", "category": "検索"},
                {"req_code": "NFR-002", "content": "画面表示性能を確保する", "category": "性能"},
                {"req_code": "FR-003", "content": "対象外の機能は unsupported とする", "category": "範囲"},
                {"req_code": "NFR-003", "content": "認証方式は未定で要相談", "category": "認証"},
            ],
        },
    )
    parsed = EvaluationBatch.model_validate(result)
    assert {item.rating for item in parsed.batch} >= {"◎", "○", "△", "×", "cần làm rõ"}
    assert len({item.note_vi for item in parsed.batch}) == 6


def test_mock_crm_solution_prefers_non_rag_direction():
    client = MockLLMClient()
    result = client.complete(
        schema=SolutionProposal,
        prompt="CRM",
        context={"request": "CRM quản lý khách hàng và lịch sử giao dịch", "intake": {"data_vi": "Excel có cấu trúc"}},
    )
    assert result["direction"] in {"none", "traditional_ml"}
    assert "RAG" in result["reason_vi"] or "rag" in result["reason_vi"].lower()


def test_mock_intake_reads_budget_and_schedule_from_rfp_attachment():
    client = MockLLMClient()
    result = client.complete(
        schema=IntakeExtraction,
        prompt="intake",
        context={
            "customer_language": "ja",
            "rfp_files": [{"filename": "ja_rfp.docx", "text": "予算目安：税別 1,500万円。2027年1月着手、9月本番展開。"}],
        },
    )
    assert result["budget_mentioned"] is True
    assert result["deadline_mentioned"] is True
    assert "1,500万円" in result["budget_note_vi"]
    assert "2027" in result["customer_schedule_note_vi"]


def test_mock_proposal_has_eight_localized_sections_and_customer_fact_tokens():
    client = MockLLMClient()
    result = client.complete(schema=ProposalOutput, prompt="proposal", context={"customer_language": "ja"})
    parsed = ProposalOutput.model_validate(result)
    assert list(parsed.sections) == [
        "introduction", "requirements_understanding", "solution_overview", "approach",
        "schedule", "team", "assumptions_and_scope", "risks",
    ]
    # Đợt 10: tiêu đề do giao diện render riêng — nội dung KHÔNG lặp tiêu đề,
    # và bản tiếng Nhật không lẫn tiếng Việt.
    for title, section in zip(
        ("はじめに", "ご要件の理解", "提案概要", "進め方", "スケジュール", "体制", "前提条件・対象外", "リスク"),
        parsed.sections.values(),
    ):
        assert not section.startswith(title)
        assert not any(ch in "àáảãạăắằẳẵặâấầẩẫậêếềểễệìíỉĩịòóỏõọôốồổỗộơớờởỡợùúủũụưứừửữựỳýỷỹỵđ" for ch in section)
    assert "{{BUDGET}}" in parsed.sections["introduction"]
    assert "{{REQUIREMENT_COUNT}}" in parsed.sections["requirements_understanding"]
