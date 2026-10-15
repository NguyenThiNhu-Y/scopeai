import re
from datetime import date
from io import BytesIO
from types import SimpleNamespace
from zipfile import ZipFile

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from app.core import db
from app.exporters.filename import content_disposition, export_filename
from app.exporters.proposal_pptx import render_slides
from app.models import Bid, ProposalDraft
from tests.helpers import put_schedule, setup_wbs_bid


def _deck_data(**overrides):
    data = {
        "project_name": "顧客管理システム（架空）",
        "customer_name": "株式会社デモ商事（架空）",
        "company_name": "サンプル会社",
        "customer_language": "ja",
        "currency": "JPY",
        "export_date": "2026-10-15",
        "sections": {"introduction": "提案内容（架空）\n- 要点1\n- 要点2", "risks": "リスク（架空）"},
        "assumptions": ["前提A"],
        "scope_out": ["対象外B"],
        "quotation": "[Presales điền]",
        "approved": False,
        "total": None,
        "wbs": [
            SimpleNamespace(task_code="1", name_vi="Thiết kế", phase="poc", parent_code=None, role=None, man_day=None),
            SimpleNamespace(task_code="1.1", name_vi="API", phase="poc", parent_code="1", role="backend", man_day=3),
            SimpleNamespace(task_code="2", name_vi="Phát triển", phase="mvp", parent_code=None, role=None, man_day=None),
            SimpleNamespace(task_code="2.1", name_vi="UI", phase="mvp", parent_code="2", role="frontend", man_day=5),
        ],
        "schedule": {
            "start_date": "2026-11-02", "end_date": "2026-12-04",
            "tasks": [
                {"task_code": "1.1", "phase": "poc", "role": "backend", "man_day": 3, "start_date": "2026-11-02", "end_date": "2026-11-04"},
                {"task_code": "2.1", "phase": "mvp", "role": "frontend", "man_day": 5, "start_date": "2026-11-16", "end_date": "2026-11-20"},
            ],
            "payment_milestones": [{"name": "Kickoff", "percent": 30, "due_date": "2026-11-02"}],
        },
        "customer_deadline": "2026-12-11",
    }
    data.update(overrides)
    return data


def _texts(slide) -> list[str]:
    return [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame and shape.text_frame.text]


def test_pptx_has_cover_agenda_sections_effort_gantt_and_gated_quotation():
    deck = Presentation(BytesIO(render_slides(_deck_data())))
    titles = [(_texts(slide) or [""])[0] for slide in deck.slides]
    assert "顧客管理システム（架空）" in " ".join(_texts(deck.slides[0]))
    assert "株式会社デモ商事（架空） 御中" in " ".join(_texts(deck.slides[0]))
    for expected in ("アジェンダ", "はじめに", "リスク", "前提条件・対象外", "工数サマリー", "実施スケジュール", "お見積り"):
        assert expected in titles
    gantt = deck.slides[titles.index("実施スケジュール")]
    bars = [shape for shape in gantt.shapes if shape.name.startswith("bar ")]
    assert len(bars) == 2
    assert bars[0].left < bars[1].left
    quotation = deck.slides[titles.index("お見積り")]
    assert "[Presales điền]" in " ".join(_texts(quotation))
    effort = deck.slides[titles.index("工数サマリー")]
    table_text = " ".join(cell.text for shape in effort.shapes if shape.has_table for row in shape.table.rows for cell in row.cells)
    assert "8.0" in table_text


def test_pptx_approved_shows_formatted_total():
    data = _deck_data(approved=True, total=2500000, quotation="2,500,000 JPY",
                      quotation_detail={"contract_model": "fixed_price", "by_phase": {"poc": 1000000, "mvp": 1500000}, "phase_man_days": {"poc": 3, "mvp": 5}, "tax_mode": "tax_excluded"})
    deck = Presentation(BytesIO(render_slides(data)))
    text = " ".join(cell.text for slide in deck.slides for shape in slide.shapes if shape.has_table for row in shape.table.rows for cell in row.cells)
    assert "2,500,000 JPY" in text and "一括請負" in " ".join(" ".join(_texts(s)) for s in deck.slides)


def test_filename_defaults_to_today_and_disposition_is_latin1_safe():
    name = export_filename("株式会社デモ", "案件", "proposal", ext="docx")
    assert name.endswith(f"_v1_{date.today().isoformat()}.docx")
    header = content_disposition(name)
    header.encode("latin-1")
    assert "filename*=UTF-8''" in header


def _bid_with_draft(client, language="vi"):
    session = db.SessionLocal()
    bid_id = setup_wbs_bid(client, session, customer_language=language)
    assert client.get(f"/api/bids/{bid_id}/proposal").status_code == 200
    put_schedule(client, bid_id)
    draft = session.query(ProposalDraft).filter_by(bid_id=bid_id).one()
    draft.assumptions_phrased = ["Giả định từ câu trả lời"]
    draft.scope_out_phrased = ["Không bao gồm vận hành"]
    session.commit()
    session.close()
    return bid_id


def test_api_exports_pptx_and_bundle_with_four_entries(client):
    bid_id = _bid_with_draft(client)
    pptx = client.get(f"/api/bids/{bid_id}/exports/proposal.pptx")
    assert pptx.status_code == 200, pptx.text
    assert pptx.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.presentationml.presentation")
    assert len(Presentation(BytesIO(pptx.content)).slides) >= 5
    bundle = client.get(f"/api/bids/{bid_id}/exports/bundle.zip")
    assert bundle.status_code == 200
    with ZipFile(BytesIO(bundle.content)) as archive:
        names = archive.namelist()
    assert [name.rsplit(".", 1)[1] for name in names] == ["docx", "xlsx", "pdf", "pptx"]
    assert all(date.today().isoformat() in name for name in names)


def test_api_docx_and_xlsx_do_not_need_pdf_converter(client, monkeypatch):
    def boom(_):
        raise AssertionError("PDF converter must not run for docx/xlsx")

    monkeypatch.setattr("app.services.export_service.docx_to_pdf", boom)
    bid_id = _bid_with_draft(client)
    docx = client.get(f"/api/bids/{bid_id}/exports/proposal.docx")
    assert docx.status_code == 200
    document = Document(BytesIO(docx.content))
    bullets = [p.text for p in document.paragraphs if p.style.name == "List Bullet"]
    assert "Giả định từ câu trả lời" in bullets and "Không bao gồm vận hành" in bullets
    assert any("Phụ lục A" in p.text for p in document.paragraphs)
    xlsx = client.get(f"/api/bids/{bid_id}/exports/estimate.xlsx")
    assert xlsx.status_code == 200
    workbook = load_workbook(BytesIO(xlsx.content))
    assert workbook["Lịch Gantt"].max_column > 6


def test_api_approved_export_renders_formatted_price(client):
    bid_id = _bid_with_draft(client, language="ja")
    session = db.SessionLocal()
    bid = session.get(Bid, bid_id)
    bid.status = "waiting_price_approval"
    session.commit()
    session.close()
    approved = client.post(
        f"/api/bids/{bid_id}/approvals/price",
        json={"decision": "approved", "contract_model": "fixed_price"},
        headers={"X-Role": "manager", "X-User": "manager"},
    )
    assert approved.status_code == 200, approved.text
    docx = client.get(f"/api/bids/{bid_id}/exports/proposal.docx")
    cells = " ".join(c.text for t in Document(BytesIO(docx.content)).tables for r in t.rows for c in r.cells)
    assert re.search(r"\d{1,3}(,\d{3})+ (VND|JPY|USD)", cells)
    workbook = load_workbook(BytesIO(client.get(f"/api/bids/{bid_id}/exports/estimate.xlsx").content))
    assert str(workbook["Tóm tắt"]["B10"].value).startswith("='Báo giá'!D")


def test_questions_export_has_instructions_and_import_reads_qa_sheet(client):
    bid_id = _bid_with_draft(client, language="ja")
    response = client.get(f"/api/bids/{bid_id}/questions/export")
    assert response.status_code == 200
    assert "_QA_v1_" in response.headers["content-disposition"]
    workbook = load_workbook(BytesIO(response.content))
    assert workbook.sheetnames == ["Q&A", "記入方法"]
    assert workbook["Q&A"].protection.sheet is True
    workbook.active = 1
    output = BytesIO()
    workbook.save(output)
    imported = client.post(
        f"/api/bids/{bid_id}/questions/import",
        files={"file": ("answers.xlsx", output.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert imported.status_code == 200, imported.text
