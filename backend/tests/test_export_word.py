from io import BytesIO

from docx import Document

from app.exporters.proposal_docx import render_proposal


def _data(locale="vi"):
    labels = {
        "introduction": "Giới thiệu",
        "solution_overview": "Tổng quan giải pháp",
        "approach": "Phương pháp",
        "risks": "Rủi ro",
        "conclusion": "Kết luận",
        "assumptions": "Giả định",
        "scope_out": "Ngoài phạm vi",
    }
    return {
        "project_name": "Dự án giả lập",
        "customer_name": "Khách hàng giả lập",
        "labels": labels,
        "sections": {"introduction": "Nội dung giả lập"},
        "assumptions": ["Giả định giả lập"],
        "scope_out": ["Ngoài phạm vi giả lập"],
        "quotation": "[Presales điền]",
    }


def test_word_opens_and_contains_cover_toc_and_sections():
    document = Document(BytesIO(render_proposal(_data())))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "Dự án giả lập" in text
    assert "Mục lục" in text
    assert "Giả định" in text
    assert "Ngoài phạm vi" in text
    assert "[Presales điền]" in text


def _full_data():
    from types import SimpleNamespace

    data = _data()
    data.update({
        "customer_language": "vi",
        "company_name": "Công ty Mẫu",
        "export_date": "2026-10-15",
        "currency": "JPY",
        "sections": {"introduction": "Mở đầu\n- Ý một\n- Ý hai", "assumptions_and_scope": "Các giả định chính"},
        "wbs": [
            SimpleNamespace(task_code="1", name_vi="Thiết kế", phase="poc", parent_code=None, role=None, man_day=None),
            SimpleNamespace(task_code="1.1", name_vi="Thiết kế API", phase="poc", parent_code="1", role="backend", man_day=2.5),
        ],
        "schedule": {
            "start_date": "2026-11-02", "end_date": "2026-11-20",
            "tasks": [{"task_code": "1.1", "phase": "poc", "role": "backend", "man_day": 2.5, "start_date": "2026-11-02", "end_date": "2026-11-04"}],
            "payment_milestones": [{"name": "Kickoff", "percent": 30, "due_date": "2026-11-02"}],
        },
    })
    return data


def test_word_has_toc_field_bullets_header_and_appendices():
    document = Document(BytesIO(render_proposal(_full_data())))
    xml = document.element.xml
    assert 'TOC \\o "1-1"' in xml
    headings = [p.text for p in document.paragraphs if p.style.name == "Heading 1"]
    assert headings[0] == "1. Giới thiệu"
    assert any(h.endswith("Phụ lục A. WBS và effort") for h in headings)
    assert any(h.endswith("Phụ lục B. Lịch tổng thể") for h in headings)
    bullets = [p.text for p in document.paragraphs if p.style.name == "List Bullet"]
    assert "Ý một" in bullets and "Giả định giả lập" in bullets and "Ngoài phạm vi giả lập" in bullets
    assert "Công ty Mẫu" in document.sections[0].header.paragraphs[0].text
    wbs_table = document.tables[0]
    assert [c.text for c in wbs_table.rows[0].cells] == ["Mã", "Công việc", "Giai đoạn", "Vai trò", "Man-day"]
    assert wbs_table.rows[2].cells[1].text == "Thiết kế API"
    assert wbs_table.rows[-1].cells[4].text == "2.5"
    cells = " ".join(c.text for t in document.tables for r in t.rows for c in r.cells)
    assert "Kickoff" in cells and "30%" in cells


def test_word_approved_quotation_is_formatted_with_currency():
    data = _full_data()
    data.update({
        "approved": True, "total": 1234567, "quotation": "1,234,567 JPY",
        "quotation_detail": {"contract_model": "fixed_price", "by_phase": {"poc": 1234567}, "phase_man_days": {"poc": 2.5}, "tax_mode": "tax_excluded"},
    })
    document = Document(BytesIO(render_proposal(data)))
    cells = " ".join(c.text for t in document.tables for r in t.rows for c in r.cells)
    assert "1,234,567 JPY" in cells
    text = "\n".join(p.text for p in document.paragraphs)
    assert "Trọn gói" in text and "[Presales điền]" not in text
