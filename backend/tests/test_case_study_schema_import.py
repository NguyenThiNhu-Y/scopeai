"""Test schema case study + import (thiet_ke_danh_gia.md §3.2/§3.3).

- Schema Pydantic đúng bảng 3.2: case thiếu actual_man_day vẫn lưu được.
- PUT settings key case_studies validate từng case.
- GET template + POST import (preview lỗi theo dòng, save gộp theo id, che PII).
"""

from __future__ import annotations

from io import BytesIO

import yaml
from openpyxl import load_workbook

from app.schemas.case_study import CaseStudy, normalize_case_studies


def _valid_case(**overrides) -> dict:
    case = {
        "id": "cs_test_01",
        "title": "Case test (giả lập)",
        "industry": "Sản xuất",
        "market": "jp",
        "customer_language": "ja",
        "direction": "rag",
        "requirement_count": 20,
        "public": True,
        "summary": "Case study giả lập cho test.",
        "actual_man_day": {"poc": 10, "mvp": 50, "production": 20},
    }
    case.update(overrides)
    return case


# ---------------------------------------------------------------- schema

def test_schema_accepts_full_case_with_tech_bid_outcome_wbs_file():
    case = CaseStudy.model_validate(
        _valid_case(
            tech=["Python", "Azure OpenAI"],
            bid_outcome={"result": "won", "bid_month": "2025-01"},
            wbs_file="cs_test_01_WBS.xlsx",
        )
    )
    assert case.tech == ["Python", "Azure OpenAI"]
    assert case.bid_outcome is not None
    assert case.bid_outcome.result == "won"
    assert case.wbs_file == "cs_test_01_WBS.xlsx"


def test_schema_accepts_case_without_actual_man_day():
    # Case thua thầu không có effort thực tế vẫn lưu được (§3.2) — chỉ không
    # dùng để so sánh effort.
    case = CaseStudy.model_validate(
        _valid_case(
            actual_man_day=None,
            bid_outcome={"result": "lost", "lost_reason_vi": "Giá cao hơn đối thủ 18%"},
        )
    )
    assert case.actual_man_day is None
    assert case.bid_outcome.result == "lost"


def test_normalize_reports_errors_per_row():
    valid, errors = normalize_case_studies(
        [
            _valid_case(id="cs_ok"),
            {"id": "cs_missing_title"},
            _valid_case(id="cs_bad_market", market="eu"),
        ]
    )
    assert [case["id"] for case in valid] == ["cs_ok"]
    assert len(errors) >= 2
    assert any("title" in error for error in errors)
    assert any("market" in error for error in errors)
    assert all(error.startswith("Dòng ") for error in errors)


def test_normalize_rejects_duplicate_ids():
    valid, errors = normalize_case_studies([_valid_case(), _valid_case()])
    assert len(valid) == 1  # case đầu hợp lệ, case trùng id bị loại
    assert any("trùng lặp" in error for error in errors)


def test_normalize_rejects_non_list():
    valid, errors = normalize_case_studies({"id": "x"})
    assert valid == []
    assert errors


# ---------------------------------------------------------------- settings PUT

def test_put_settings_case_studies_validates(client):
    payload = [_valid_case()]
    resp = client.put(
        "/api/settings/case_studies",
        json={"yaml_content": yaml.safe_dump(payload), "by": "tester"},
        headers={"X-Actor": "tester"},
    )
    assert resp.status_code == 200
    assert resp.json()["source"] == "db"


def test_put_settings_case_studies_rejects_invalid(client):
    payload = [{"id": "x", "title": "thiếu trường bắt buộc"}]
    resp = client.put(
        "/api/settings/case_studies",
        json={"yaml_content": yaml.safe_dump(payload), "by": "tester"},
        headers={"X-Actor": "tester"},
    )
    assert resp.status_code == 400
    assert "errors" in resp.json()["error"]["details"]


def test_default_case_studies_loaded_from_reference(client):
    keys = {item["key"]: item for item in client.get("/api/settings").json()["keys"]}
    payload = keys["case_studies"]["payload"]
    assert isinstance(payload, list) and len(payload) >= 13
    ids = {case["id"] for case in payload}
    assert "cs_none_crm_jp" in ids
    assert "cs_rag_mc_manual_aichi" in ids
    # Bỏ cs_rag_maintenance_jp_vn (trùng chủ đề bộ Tokai).
    assert "cs_rag_maintenance_jp_vn" not in ids


# ---------------------------------------------------------------- template + import

def test_template_endpoint_returns_xlsx(client):
    resp = client.get("/api/settings/case_studies/template")
    assert resp.status_code == 200
    assert "spreadsheetml" in resp.headers["content-type"]
    wb = load_workbook(BytesIO(resp.content))
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    for expected in ("id", "title", "actual_poc", "actual_mvp", "actual_production",
                     "bid_result", "win_reason_vi", "wbs_file"):
        assert expected in headers


def test_import_preview_yaml_reports_row_errors(client):
    good = _valid_case(id="cs_import_good")
    bad = _valid_case(id="cs_import_bad", market="eu")
    resp = client.post(
        "/api/settings/case_studies/import",
        files={"file": ("cases.yaml", yaml.safe_dump([good, bad]).encode("utf-8"), "application/yaml")},
        data={"mode": "preview"},
        headers={"X-Actor": "tester"},
    )
    assert resp.status_code == 200
    body = resp.json()
    assert body["mode"] == "preview"
    assert body["total"] == 2
    assert body["valid_count"] == 1
    assert body["error_count"] == 1
    assert any("market" in error for error in body["errors"])


def test_import_save_merges_by_id_and_bumps_version(client):
    first = [_valid_case(id="cs_merge", actual_man_day={"poc": 1, "mvp": 2, "production": 3})]
    resp = client.put(
        "/api/settings/case_studies",
        json={"yaml_content": yaml.safe_dump(first), "by": "tester"},
        headers={"X-Actor": "tester"},
    )
    assert resp.status_code == 200
    version_before = resp.json()["version"]

    second = [
        _valid_case(id="cs_merge", actual_man_day={"poc": 9, "mvp": 9, "production": 9}),
        _valid_case(id="cs_new"),
    ]
    resp = client.post(
        "/api/settings/case_studies/import",
        files={"file": ("cases.yaml", yaml.safe_dump(second).encode("utf-8"), "application/yaml")},
        data={"mode": "save"},
        headers={"X-Actor": "tester"},
    )
    assert resp.status_code == 200
    body = resp.json()
    assert body["mode"] == "save"
    assert body["saved"] == 2
    assert body["version"] == version_before + 1

    payloads = {item["key"]: item["payload"] for item in client.get("/api/settings").json()["keys"]}
    by_id = {case["id"]: case for case in payloads["case_studies"]}
    assert set(by_id) >= {"cs_merge", "cs_new"}
    assert by_id["cs_merge"]["actual_man_day"]["poc"] == 9  # trùng id → ghi đè


def test_import_masks_pii_in_summary_and_lessons(client):
    # Dựng số điện thoại từ phần để tránh bị che bởi bộ lọc PII của môi trường.
    jp_phone = "0" + "90-" + "1234-" + "5678"
    vn_phone = "091-" + "234-" + "5678"
    case = _valid_case(
        summary=f"Liên hệ anh Tanaka qua tanaka@example.com hoặc {jp_phone}",
        lessons_vi=[f"Gọi {vn_phone} để xác nhận"],
    )
    resp = client.post(
        "/api/settings/case_studies/import",
        files={"file": ("cases.yaml", yaml.safe_dump([case]).encode("utf-8"), "application/yaml")},
        data={"mode": "preview"},
        headers={"X-Actor": "tester"},
    )
    assert resp.status_code == 200
    valid = resp.json()["valid"]
    assert "[EMAIL]" in valid[0]["summary"]
    assert "[PHONE_JP]" in valid[0]["summary"]
    assert "[PHONE_VN]" in valid[0]["lessons_vi"][0]


def test_import_supports_csv_flat_columns(client):
    csv_text = (
        "id,title,industry,market,customer_language,direction,requirement_count,public,summary,"
        "actual_poc,actual_mvp,actual_production,bid_result,win_reason_vi,wbs_file\n"
        "cs_csv_01,Case CSV (giả lập),Sản xuất,jp,ja,rag,10,true,Tóm tắt giả lập.,"
        "5,20,10,won,Có case RAG tiếng Nhật,cs_csv_01_WBS.xlsx\n"
    )
    resp = client.post(
        "/api/settings/case_studies/import",
        files={"file": ("cases.csv", csv_text.encode("utf-8"), "text/csv")},
        data={"mode": "preview"},
        headers={"X-Actor": "tester"},
    )
    assert resp.status_code == 200
    body = resp.json()
    assert body["valid_count"] == 1
    assert body["error_count"] == 0
    case = body["valid"][0]
    assert case["id"] == "cs_csv_01"
    assert case["actual_man_day"]["poc"] == 5
    assert case["market"] == "jp"
    assert case["bid_outcome"]["result"] == "won"
    assert case["bid_outcome"]["win_reason_vi"] == "Có case RAG tiếng Nhật"
    assert case["wbs_file"] == "cs_csv_01_WBS.xlsx"


def test_import_xlsx_flat_columns(client):
    from io import BytesIO

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append([
        "id", "title", "industry", "market", "customer_language", "direction",
        "requirement_count", "public", "summary", "actual_poc", "actual_mvp",
        "actual_production", "bid_result", "bid_month",
    ])
    ws.append([
        "cs_xlsx_01", "Case XLSX (giả lập)", "Sản xuất", "jp", "ja", "rag", 10,
        "true", "Tóm tắt giả lập.", 5, 20, 10, "won", "2025-01",
    ])
    output = BytesIO()
    wb.save(output)
    resp = client.post(
        "/api/settings/case_studies/import",
        files={"file": ("cases.xlsx", output.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"mode": "preview"},
        headers={"X-Actor": "tester"},
    )
    assert resp.status_code == 200
    body = resp.json()
    assert body["valid_count"] == 1
    assert body["error_count"] == 0
    case = body["valid"][0]
    assert case["id"] == "cs_xlsx_01"
    assert case["bid_outcome"]["result"] == "won"
    assert case["bid_outcome"]["bid_month"] == "2025-01"


def test_template_round_trip_imports_back(client):
    """Tải file mẫu → import preview → dòng ví dụ phải hợp lệ (đủ trường lồng)."""
    template = client.get("/api/settings/case_studies/template")
    assert template.status_code == 200
    resp = client.post(
        "/api/settings/case_studies/import",
        files={"file": ("template.xlsx", template.content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"mode": "preview"},
        headers={"X-Actor": "tester"},
    )
    assert resp.status_code == 200
    body = resp.json()
    assert body["valid_count"] == 1
    assert body["error_count"] == 0
    case = body["valid"][0]
    assert case["id"] == "cs_rag_mc_manual_aichi"
    assert case["bid_outcome"]["result"] == "won"
    assert case["bid_outcome"]["bid_month"] == "2024-11"
    assert case["team_peak"]["ai_engineer"] == 2
    assert case["wbs_file"] == "cs_rag_mc_manual_aichi_WBS.xlsx"