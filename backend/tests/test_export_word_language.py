from io import BytesIO

from docx import Document

from app.exporters.proposal_docx import render_proposal


def test_word_japanese_labels_are_rendered():
    data = {
        "project_name": "顧客管理システム（架空）",
        "customer_name": "株式会社デモ商事（仮名・架空）",
        "labels": {"introduction": "はじめに", "assumptions": "前提条件", "scope_out": "スコープ外"},
        "sections": {"introduction": "提案内容（架空）"},
        "assumptions": ["前提の説明"],
        "scope_out": ["対象外の説明"],
        "quotation": "[Presales điền]",
    }
    document = Document(BytesIO(render_proposal(data)))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "はじめに" in text
    assert "前提条件" in text
    assert "スコープ外" in text


def test_word_english_labels_are_rendered():
    data = {
        "project_name": "Fictional CRM Project",
        "customer_name": "Fictional Customer",
        "labels": {"introduction": "Introduction", "assumptions": "Assumptions", "scope_out": "Scope out"},
        "sections": {"introduction": "Fictional proposal"},
        "assumptions": ["Fictional assumption"],
        "scope_out": ["Fictional exclusion"],
        "quotation": "[Presales điền]",
    }
    document = Document(BytesIO(render_proposal(data)))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "Introduction" in text
    assert "Assumptions" in text
    assert "Scope out" in text
