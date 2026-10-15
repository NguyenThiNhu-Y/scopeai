"""Proposal summary deck (.pptx) — every number comes from persisted, code-computed data."""

from __future__ import annotations

from datetime import date, timedelta
from io import BytesIO

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from app.exporters.common import (
    CONTRACT_LABELS,
    PHASES,
    PRICE_PLACEHOLDER,
    as_date,
    format_money,
    labels_for,
    leaf_man_days,
    phase_label,
    role_label,
    wbs_rows,
)
from app.exporters.proposal_docx import outline

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
GREY = RGBColor(0x59, 0x59, 0x59)
LIGHT = RGBColor(0xD9, 0xE2, 0xF3)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RED = RGBColor(0xC0, 0x00, 0x00)
PHASE_COLORS = {"poc": RGBColor(0x8F, 0xAA, 0xDC), "mvp": RGBColor(0xA9, 0xD1, 0x8E), "production": RGBColor(0xF4, 0xB1, 0x83)}
SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)
MAX_GANTT_ROWS = 14


class _Deck:
    def __init__(self, data: dict):
        self.data = data
        self.language = data.get("customer_language") or "vi"
        self.labels = labels_for(self.language, data.get("labels"))
        self.font = "Yu Gothic" if self.language == "ja" else "Arial"
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = SLIDE_W, SLIDE_H
        self.prs.core_properties.title = data["project_name"]
        self.prs.core_properties.author = data.get("company_name") or ""

    def style(self, run, size: float, bold: bool = False, color: RGBColor = GREY) -> None:
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        run.font.name = self.font
        rpr = run._r.get_or_add_rPr()
        for tag in ("a:ea", "a:cs"):
            element = rpr.find(qn(tag))
            if element is None:
                element = rpr.makeelement(qn(tag), {})
                rpr.append(element)
            element.set("typeface", self.font)

    def text(self, slide, left, top, width, height, lines: list[str], size: float = 16, bold: bool = False,
             color: RGBColor = GREY, bullet: bool = False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
        box = slide.shapes.add_textbox(left, top, width, height)
        frame = box.text_frame
        frame.word_wrap = True
        frame.vertical_anchor = anchor
        for index, line in enumerate(lines):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.alignment = align
            paragraph.space_after = Pt(6)
            run = paragraph.add_run()
            run.text = f"• {line}" if bullet else line
            self.style(run, size, bold, color)
        return box

    def rect(self, slide, left, top, width, height, color: RGBColor, shape=MSO_SHAPE.RECTANGLE):
        item = slide.shapes.add_shape(shape, left, top, width, height)
        item.fill.solid()
        item.fill.fore_color.rgb = color
        item.line.fill.background()
        item.shadow.inherit = False
        return item

    def slide(self, title: str):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.rect(slide, 0, 0, SLIDE_W, Inches(1.0), NAVY)
        self.text(slide, Inches(0.5), Inches(0.15), Inches(12.3), Inches(0.75), [title], 26, True, WHITE, anchor=MSO_ANCHOR.MIDDLE)
        footer = " | ".join(x for x in (self.data.get("company_name"), self.data["project_name"]) if x)
        self.text(slide, Inches(0.5), Inches(7.0), Inches(10), Inches(0.35), [footer], 10)
        self.text(slide, Inches(11.8), Inches(7.0), Inches(1.0), Inches(0.35), [str(len(self.prs.slides))], 10, align=PP_ALIGN.RIGHT)
        return slide

    def table(self, slide, headers: list[str], rows: list[list[str]], left, top, col_widths_in: list[float], bold_last: bool = False):
        shape = slide.shapes.add_table(len(rows) + 1, len(headers), left, top, Inches(sum(col_widths_in)), Inches(0.4 * (len(rows) + 1)))
        table = shape.table
        for index, width in enumerate(col_widths_in):
            table.columns[index].width = Inches(width)
        for r, values in enumerate([headers] + rows):
            for c, value in enumerate(values):
                cell = table.cell(r, c)
                cell.text = ""
                paragraph = cell.text_frame.paragraphs[0]
                paragraph.alignment = PP_ALIGN.RIGHT if c > 0 and r > 0 and _looks_numeric(value) else PP_ALIGN.LEFT
                header = r == 0
                run = paragraph.add_run()
                run.text = str(value)
                self.style(run, 13, header or (bold_last and r == len(rows)), WHITE if header else GREY)
                cell.fill.solid()
                cell.fill.fore_color.rgb = NAVY if header else (LIGHT if bold_last and r == len(rows) else WHITE)
        return table

    # --- slides -------------------------------------------------------------

    def cover(self) -> None:
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.rect(slide, 0, 0, Inches(0.35), SLIDE_H, NAVY)
        self.text(slide, Inches(1.0), Inches(2.0), Inches(11.3), Inches(1.6), [self.data["project_name"]], 36, True, NAVY)
        self.text(slide, Inches(1.0), Inches(3.5), Inches(11.3), Inches(0.6), [self.labels["proposal"]], 20)
        customer = f"{self.labels['customer_prefix']}{self.data['customer_name']}{self.labels['customer_suffix']}"
        lines = [customer]
        if self.data.get("export_date"):
            lines.append(f"{self.labels['date']}: {self.data['export_date']}")
        if self.data.get("company_name"):
            lines.append(f"{self.labels['prepared_by']}: {self.data['company_name']}")
        self.text(slide, Inches(1.0), Inches(5.0), Inches(11.3), Inches(1.8), lines, 16)

    def agenda(self) -> None:
        slide = self.slide(self.labels["agenda"])
        self.text(slide, Inches(0.8), Inches(1.4), Inches(11.7), Inches(5.4), outline(self.data, self.labels), 20)

    def sections(self) -> None:
        limit = 420 if self.language == "ja" else 750
        for key, value in (self.data.get("sections") or {}).items():
            lines = [line.strip().lstrip("-*•・ ").strip() for line in str(value or "").splitlines() if line.strip()]
            chunks, current, size = [], [], 0
            for line in lines:
                if current and (size + len(line) > limit or len(current) >= 9):
                    chunks.append(current)
                    current, size = [], 0
                current.append(line)
                size += len(line)
            if current:
                chunks.append(current)
            for index, chunk in enumerate(chunks or [[""]]):
                title = self.labels.get(key, key) + (f" {self.labels['continued']}" if index else "")
                slide = self.slide(title)
                self.text(slide, Inches(0.8), Inches(1.4), Inches(11.7), Inches(5.4), chunk, 18, bullet=len(chunk) > 1)
            if key == "assumptions_and_scope":
                self.assumptions()

    def assumptions(self) -> None:
        if not (self.data.get("assumptions") or self.data.get("scope_out")):
            return
        slide = self.slide(self.labels["assumptions_and_scope"])
        for offset, key in enumerate(("assumptions", "scope_out")):
            left = Inches(0.8 + offset * 6.1)
            self.text(slide, left, Inches(1.3), Inches(5.8), Inches(0.5), [self.labels[key]], 20, True, NAVY)
            items = [str(item) for item in self.data.get(key) or []][:10]
            self.text(slide, left, Inches(1.9), Inches(5.8), Inches(4.9), items, 14, bullet=True)

    def effort(self) -> None:
        items = list(self.data.get("wbs") or [])
        if not items:
            return
        leaves = [item for item, _, is_parent in wbs_rows(items) if not is_parent]
        slide = self.slide(self.labels["effort_summary"])
        by_phase = [[phase_label(p, self.language), _md(sum(float(i.man_day or 0) for i in leaves if getattr(i, "phase", None) == p))]
                    for p in PHASES if any(getattr(i, "phase", None) == p for i in leaves)]
        by_phase.append([self.labels["total"], _md(leaf_man_days(items))])
        self.table(slide, [self.labels["phase"], self.labels["man_day"]], by_phase, Inches(0.8), Inches(1.5), [3.2, 2.2], bold_last=True)
        roles: dict[str, float] = {}
        for item in leaves:
            roles[role_label(getattr(item, "role", None))] = roles.get(role_label(getattr(item, "role", None)), 0.0) + float(item.man_day or 0)
        role_rows = [[role, _md(value)] for role, value in roles.items()][:11]
        self.table(slide, [self.labels["role"], self.labels["man_day"]], role_rows, Inches(7.0), Inches(1.5), [3.2, 2.2])

    def gantt(self) -> None:
        schedule = self.data.get("schedule") or {}
        tasks = [t for t in schedule.get("tasks") or [] if as_date(t.get("start_date")) and as_date(t.get("end_date"))]
        if not tasks:
            return
        slide = self.slide(self.labels["gantt"])
        rows = _gantt_rows(tasks, self.data.get("wbs") or [], self.language)
        milestones = [(m.get("name", ""), as_date(m.get("due_date"))) for m in schedule.get("payment_milestones") or []]
        milestones = [(name, day) for name, day in milestones if day]
        deadline = as_date(self.data.get("customer_deadline"))
        days = [r[1] for r in rows] + [r[2] for r in rows] + [d for _, d in milestones] + ([deadline] if deadline else [])
        first = min(days) - timedelta(days=min(days).weekday())
        last = max(days) + timedelta(days=1)
        span = max((last - first).days, 1)
        left, label_w, top = Inches(0.4), Inches(3.4), Inches(1.75)
        chart_left, chart_w = left + label_w, Inches(12.5) - label_w
        row_h = min(Inches(0.36), int((Inches(5.0)) / max(len(rows) + 1, 1)))

        def x(day: date) -> int:
            return int(chart_left + chart_w * ((day - first).days / span))

        step = 7 if span <= 7 * 13 else 28 if span <= 28 * 13 else 91
        tick = first
        while tick <= last:
            self.text(slide, x(tick), Inches(1.2), Inches(1.0), Inches(0.35), [f"{tick:%m/%d}"], 10)
            line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x(tick), Inches(1.55), x(tick), top + row_h * (len(rows) + 1))
            line.line.color.rgb = LIGHT
            tick += timedelta(days=step)
        for index, (label, start, end, phase) in enumerate(rows):
            y = top + row_h * index
            self.text(slide, left, y, label_w, row_h, [label], 11, anchor=MSO_ANCHOR.MIDDLE)
            bar = self.rect(slide, x(start), int(y + row_h * 0.2), max(x(end + timedelta(days=1)) - x(start), Inches(0.05)), int(row_h * 0.6), PHASE_COLORS.get(phase, LIGHT))
            bar.name = f"bar {label}"
        y = top + row_h * len(rows)
        if milestones:
            self.text(slide, left, y, label_w, row_h, [self.labels["payment_milestones"]], 11, True, anchor=MSO_ANCHOR.MIDDLE)
            size = Inches(0.2)
            for name, day in milestones:
                self.rect(slide, x(day) - size // 2, int(y + (row_h - size) / 2), size, size, NAVY, MSO_SHAPE.DIAMOND)
                self.text(slide, x(day) - Inches(0.6), int(y + row_h * 0.8), Inches(1.4), Inches(0.5), [name], 9, align=PP_ALIGN.CENTER)
        if deadline:
            mark = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x(deadline), Inches(1.55), x(deadline), y + row_h)
            mark.line.color.rgb = RED
            mark.line.width = Pt(2)
            self.text(slide, x(deadline) - Inches(0.9), Inches(6.5), Inches(1.8), Inches(0.4), [f"{self.labels['customer_deadline']} {deadline:%Y-%m-%d}"], 10, True, RED, align=PP_ALIGN.CENTER)

    def quotation(self) -> None:
        slide = self.slide(self.labels["quotation"])
        detail = self.data.get("quotation_detail") if self.data.get("approved") else None
        currency = self.data.get("currency")
        if not detail:
            self.text(slide, Inches(0.8), Inches(2.4), Inches(11.7), Inches(1.0), [self.data.get("quotation") or PRICE_PLACEHOLDER], 32, True, NAVY, align=PP_ALIGN.CENTER)
            self.text(slide, Inches(0.8), Inches(3.5), Inches(11.7), Inches(0.6), [self.labels["price_pending"]], 16, align=PP_ALIGN.CENTER)
        else:
            by_phase = detail.get("by_phase") or {}
            days = detail.get("phase_man_days") or {}
            rows = [[phase_label(p, self.language), _md(days.get(p)), format_money(by_phase[p], currency)] for p in PHASES if p in by_phase]
            rows.append([self.labels["total"], _md(sum(float(v or 0) for v in days.values())), format_money(self.data.get("total"), currency)])
            self.table(slide, [self.labels["phase"], self.labels["man_day"], self.labels["amount"]], rows, Inches(0.8), Inches(1.5), [3.2, 2.0, 3.4], bold_last=True)
            model = CONTRACT_LABELS.get(self.language, CONTRACT_LABELS["vi"]).get(detail.get("contract_model"), detail.get("contract_model"))
            notes = [f"{self.labels['contract_model']}: {model}",
                     self.labels["tax_included"] if detail.get("tax_mode") == "tax_included" else self.labels["tax_excluded"]]
            if detail.get("valid_until"):
                notes.append(f"{self.labels['valid_until']}: {detail['valid_until']}")
            self.text(slide, Inches(9.0), Inches(1.5), Inches(3.8), Inches(2.5), notes, 14)
        milestones = (self.data.get("schedule") or {}).get("payment_milestones") or []
        if milestones:
            rows = [[m.get("name", ""), f"{float(m.get('percent') or 0):g}%", str(m.get("due_date") or "")] for m in milestones][:6]
            self.table(slide, [self.labels["milestone"], self.labels["percent"], self.labels["due_date"]], rows, Inches(0.8), Inches(4.3), [4.2, 1.6, 2.8])

    def render(self) -> bytes:
        self.cover()
        self.agenda()
        self.sections()
        if "assumptions_and_scope" not in (self.data.get("sections") or {}):
            self.assumptions()
        self.effort()
        self.gantt()
        self.quotation()
        stream = BytesIO()
        self.prs.save(stream)
        return stream.getvalue()


def _md(value) -> str:
    return f"{float(value or 0):,.1f}"


def _looks_numeric(value) -> bool:
    text = str(value).replace(",", "").replace("%", "").split(" ")[0]
    try:
        float(text)
        return True
    except ValueError:
        return False


def _gantt_rows(tasks: list[dict], wbs, language: str) -> list[tuple[str, date, date, str]]:
    """Group leaf tasks by WBS parent; fall back to one row per phase when there are too many groups."""
    by_code = {item.task_code: item for item in wbs}
    groups: dict[str, list[dict]] = {}
    for task in tasks:
        parent = getattr(by_code.get(task.get("task_code")), "parent_code", None) or task.get("task_code")
        groups.setdefault(parent, []).append(task)
    if len(groups) > MAX_GANTT_ROWS:
        groups = {}
        for task in tasks:
            groups.setdefault(task.get("phase") or "", []).append(task)
        names = {phase: phase_label(phase, language) for phase in groups}
    else:
        names = {code: f"{code} {getattr(by_code.get(code), 'name_vi', '')}".strip() for code in groups}
    rows = []
    for key, items in groups.items():
        start = min(as_date(t["start_date"]) for t in items)
        end = max(as_date(t["end_date"]) for t in items)
        rows.append((names[key][:40], start, end, items[0].get("phase") or ""))
    return sorted(rows, key=lambda r: (r[1], r[0]))


def render_slides(data: dict) -> bytes:
    return _Deck(data).render()
