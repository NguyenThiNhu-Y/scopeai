from __future__ import annotations

from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell

from app.models import Attachment, Evaluation, Requirement


RATING_LABELS = {"◎": "◎", "○": "○", "△": "△", "×": "×"}


def _find_header(values, aliases):
    for index, value in enumerate(values, 1):
        text = str(value or "").strip().lower()
        if any(alias in text for alias in aliases):
            return index
    return None


def render_requirement_workbook(
    session, bid_id: int, customer_language: str, note_language: str | None = None
) -> tuple[str, bytes]:
    attachment = session.query(Attachment).filter_by(bid_id=bid_id, file_type="requirement").order_by(Attachment.id).first()
    if attachment is None:
        raise ValueError("Không tìm thấy file Excel requirement gốc")
    source = Path(attachment.original_path or attachment.path)
    if source.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("Chỉ hỗ trợ xuất lại file requirement dạng XLSX/XLSM")
    workbook = load_workbook(source, keep_vba=source.suffix.lower() == ".xlsm")
    requirements = session.query(Requirement).filter_by(bid_id=bid_id).all()
    evaluation_by_req = {e.requirement_id: e for e in session.query(Evaluation).filter_by(bid_id=bid_id).all()}
    language = note_language or customer_language
    note_prefix = {"ja": "対応", "en": "Coverage", "vi": "Đáp ứng"}.get(language, "Coverage")
    for requirement in requirements:
        if requirement.source_sheet not in workbook.sheetnames:
            continue
        sheet = workbook[requirement.source_sheet]
        header_row = int((requirement.source_cols or {}).get("__header_row__", 1))
        headers = [sheet.cell(header_row, column).value for column in range(1, sheet.max_column + 1)]
        response_column = _find_header(headers, ("対応可否", "coverage", "đáp ứng"))
        note_column = _find_header(headers, ("補足説明", "補足", "note", "ghi chú"))
        if response_column is None:
            response_column = sheet.max_column + 1
            cell = sheet.cell(header_row, response_column)
            if not isinstance(cell, MergedCell):
                cell.value = "対応可否"
        if note_column is None:
            note_column = sheet.max_column + 1
            cell = sheet.cell(header_row, note_column)
            if not isinstance(cell, MergedCell):
                cell.value = "補足説明"
        evaluation = evaluation_by_req.get(requirement.id)
        if evaluation is None:
            rating, note = "要確認", f"{note_prefix}: chưa có kết quả đánh giá"
        else:
            rating, note = RATING_LABELS.get(evaluation.rating, "要確認"), evaluation.note_vi
        row = requirement.source_row
        response_cell = sheet.cell(row, response_column)
        if not isinstance(response_cell, MergedCell):
            response_cell.value = rating
        note_cell = sheet.cell(row, note_column)
        if not isinstance(note_cell, MergedCell):
            note_cell.value = note
    output = BytesIO()
    workbook.save(output)
    return source.name, output.getvalue()
