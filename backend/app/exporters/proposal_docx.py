from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor

from app.exporters.common import (
    CONTRACT_LABELS,
    PHASES,
    PRICE_PLACEHOLDER,
    format_money,
    labels_for,
    leaf_man_days,
    phase_label,
    phase_summary,
    role_label,
    wbs_rows,
)

BULLET_PREFIXES = ("- ", "* ", "• ", "・")
HEADER_FILL = "D9E2F3"
PARENT_FILL = "F2F2F2"


def _set_style_font(style, latin: str, east_asia: str, size: float | None = None, color: RGBColor | None = None) -> None:
    style.font.name = latin
    if size:
        style.font.size = Pt(size)
    if color is not None:
        style.font.color.rgb = color
    rfonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        rfonts.attrib.pop(qn(f"w:{attr}"), None)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), east_asia)


def _setup_styles(document, language: str) -> None:
    latin, east_asia = ("Yu Gothic", "Yu Gothic") if language == "ja" else ("Arial", "Arial")
    navy = RGBColor(0x1F, 0x3A, 0x5F)
    _set_style_font(document.styles["Normal"], latin, east_asia, 10.5)
    _set_style_font(document.styles["Title"], latin, east_asia, 26, navy)
    _set_style_font(document.styles["Heading 1"], latin, east_asia, 15, navy)
    _set_style_font(document.styles["Heading 2"], latin, east_asia, 12, navy)
    _set_style_font(document.styles["List Bullet"], latin, east_asia, 10.5)
    document.styles["Normal"].paragraph_format.space_after = Pt(4)


def _add_field(paragraph, instruction: str, cached_lines: list[str]) -> None:
    """Complex field with a pre-rendered cached result, so viewers that never update fields still show it."""
    def fld(kind: str) -> None:
        element = OxmlElement("w:fldChar")
        element.set(qn("w:fldCharType"), kind)
        paragraph.add_run()._r.append(element)

    fld("begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    paragraph.add_run()._r.append(instr)
    fld("separate")
    for index, line in enumerate(cached_lines):
        run = paragraph.add_run(line)
        if index < len(cached_lines) - 1:
            run.add_break(WD_BREAK.LINE)
    fld("end")


def _shade(cell, fill: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def _table(document, headers: list[str], rows: list[list], widths_cm: list[float], bold_rows: set[int] = frozenset(),
           numeric_cols: set[int] = frozenset(), indents: dict[int, int] | None = None):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        run = cell.paragraphs[0].add_run(str(header))
        run.bold = True
        run.font.size = Pt(9)
        _shade(cell, HEADER_FILL)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for col, value in enumerate(values):
            paragraph = cells[col].paragraphs[0]
            run = paragraph.add_run("" if value is None else str(value))
            run.font.size = Pt(9)
            run.bold = row_index in bold_rows
            if col in numeric_cols:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if col == 1 and indents and indents.get(row_index):
                paragraph.paragraph_format.left_indent = Cm(0.35 * indents[row_index])
            if row_index in bold_rows:
                _shade(cells[col], PARENT_FILL)
    for row in table.rows:
        for index, width in enumerate(widths_cm):
            row.cells[index].width = Cm(width)
    document.add_paragraph()
    return table


def _add_text(document, value: str) -> None:
    for raw in str(value or "").splitlines():
        line = raw.strip()
        if not line:
            continue
        prefix = next((p for p in BULLET_PREFIXES if line.startswith(p)), None)
        if prefix:
            document.add_paragraph(line[len(prefix):].strip(), style="List Bullet")
        else:
            document.add_paragraph(line)


def _num(value) -> str:
    return f"{float(value or 0):,.1f}"


def _cover(document, data: dict, labels: dict) -> None:
    for _ in range(6):
        document.add_paragraph()
    document.add_heading(data["project_name"], 0)
    subtitle = document.add_paragraph()
    run = subtitle.add_run(labels["proposal"])
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor(0x40, 0x40, 0x40)
    for _ in range(4):
        document.add_paragraph()
    customer = document.add_paragraph()
    customer_run = customer.add_run(f"{labels['customer_prefix']}{data['customer_name']}{labels['customer_suffix']}")
    customer_run.font.size = Pt(14)
    customer_run.bold = True
    if data.get("export_date"):
        document.add_paragraph(f"{labels['date']}: {data['export_date']}")
    if data.get("company_name"):
        document.add_paragraph(f"{labels['prepared_by']}: {data['company_name']}")


def _header_footer(section, data: dict, labels: dict) -> None:
    section.different_first_page_header_footer = True
    header = section.header.paragraphs[0]
    header.text = " | ".join(x for x in (data.get("company_name"), data["project_name"]) if x)
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0x70, 0x70, 0x70)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run(f"{labels['page']} ")
    _add_field(footer, "PAGE", ["1"])


def _assumptions(document, data: dict, labels: dict) -> None:
    for key in ("assumptions", "scope_out"):
        items = data.get(key) or []
        if not items:
            continue
        document.add_heading(labels[key], 2)
        for item in items:
            document.add_paragraph(str(item), style="List Bullet")


def _quotation(document, data: dict, labels: dict, language: str) -> None:
    detail = data.get("quotation_detail") if data.get("approved") else None
    currency = data.get("currency")
    if not detail:
        document.add_paragraph(data.get("quotation") or PRICE_PLACEHOLDER)
        return
    model = CONTRACT_LABELS.get(language, CONTRACT_LABELS["vi"]).get(detail.get("contract_model"), detail.get("contract_model"))
    document.add_paragraph(f"{labels['contract_model']}: {model}")
    document.add_paragraph(f"{labels['currency']}: {currency}")
    by_phase = detail.get("by_phase") or {}
    phase_days = detail.get("phase_man_days") or {}
    rows = [[phase_label(p, language), _num(phase_days.get(p)), format_money(by_phase[p], currency)] for p in PHASES if p in by_phase]
    rows.append([labels["total"], _num(sum(float(v or 0) for v in phase_days.values())), format_money(data.get("total"), currency)])
    _table(document, [labels["phase"], labels["man_day"], labels["amount"]], rows, [6, 3.5, 6], bold_rows={len(rows) - 1}, numeric_cols={1, 2})
    document.add_paragraph(labels["tax_included"] if detail.get("tax_mode") == "tax_included" else labels["tax_excluded"])
    if detail.get("valid_until"):
        document.add_paragraph(f"{labels['valid_until']}: {detail['valid_until']}")


def _wbs_appendix(document, data: dict, labels: dict, language: str) -> None:
    items = list(data.get("wbs") or [])
    rows, bold, indents = [], set(), {}
    for index, (item, level, is_parent) in enumerate(wbs_rows(items)):
        rows.append([
            item.task_code,
            getattr(item, "name_vi", ""),
            phase_label(getattr(item, "phase", None), language),
            "" if is_parent else role_label(getattr(item, "role", None)),
            "" if is_parent else _num(item.man_day),
        ])
        if is_parent:
            bold.add(index)
        indents[index] = level - 1
    rows.append([labels["total"], "", "", "", _num(leaf_man_days(items))])
    bold.add(len(rows) - 1)
    _table(document, [labels["task_code"], labels["task_name"], labels["phase"], labels["role"], labels["man_day"]],
           rows, [2.2, 7.3, 2.4, 2.4, 2.2], bold, {4}, indents)


def _schedule_appendix(document, data: dict, labels: dict, language: str) -> None:
    schedule = data.get("schedule") or {}
    tasks = schedule.get("tasks") or []
    if not tasks:
        document.add_paragraph(labels["no_schedule"])
        return
    document.add_paragraph(f"{labels['start']}: {schedule.get('start_date') or ''}    {labels['end']}: {schedule.get('end_date') or ''}")
    rows = [[phase_label(r["phase"], language), r["start"].isoformat(), r["end"].isoformat(), _num(r["man_day"])] for r in phase_summary(tasks)]
    _table(document, [labels["phase"], labels["start"], labels["end"], labels["man_day"]], rows, [4.5, 4, 4, 3], numeric_cols={3})
    milestones = schedule.get("payment_milestones") or []
    if milestones:
        document.add_heading(labels["payment_milestones"], 2)
        rows = [[m.get("name", ""), f"{float(m.get('percent') or 0):g}%", m.get("due_date", "")] for m in milestones]
        _table(document, [labels["milestone"], labels["percent"], labels["due_date"]], rows, [7, 3, 5], numeric_cols={1})
    if data.get("customer_deadline"):
        document.add_paragraph(f"{labels['customer_deadline']}: {data['customer_deadline']}")


def outline(data: dict, labels: dict) -> list[str]:
    """Numbered top-level headings, shared by the Word TOC and the slide agenda."""
    sections = data.get("sections") or {}
    headings = [labels.get(key, key) for key in sections]
    if (data.get("assumptions") or data.get("scope_out")) and "assumptions_and_scope" not in sections:
        headings.append(labels["assumptions_and_scope"])
    headings.append(labels["quotation"])
    if data.get("wbs"):
        headings.append(labels["wbs_appendix"])
    if data.get("schedule"):
        headings.append(labels["schedule_appendix"])
    return [f"{index}. {title}" for index, title in enumerate(headings, 1)]


def render_proposal(data: dict) -> bytes:
    language = data.get("customer_language") or "vi"
    labels = labels_for(language, data.get("labels"))
    document = Document()
    document.core_properties.title = data["project_name"]
    document.core_properties.author = data.get("company_name") or ""
    _setup_styles(document, language)
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.left_margin = section.right_margin = Cm(2.2)
    section.top_margin = section.bottom_margin = Cm(2.0)
    _header_footer(section, data, labels)

    _cover(document, data, labels)
    document.add_page_break()

    sections = dict(data.get("sections") or {})
    numbered = outline(data, labels)
    toc_title = document.add_paragraph()
    toc_run = toc_title.add_run(labels["toc"])
    toc_run.bold = True
    toc_run.font.size = Pt(15)
    _add_field(document.add_paragraph(), 'TOC \\o "1-1" \\h \\z \\u', numbered)
    document.add_page_break()

    counter = iter(numbered)
    for key, value in sections.items():
        document.add_heading(next(counter), 1)
        _add_text(document, value)
        if key == "assumptions_and_scope":
            _assumptions(document, data, labels)
    if (data.get("assumptions") or data.get("scope_out")) and "assumptions_and_scope" not in sections:
        document.add_heading(next(counter), 1)
        _assumptions(document, data, labels)
    document.add_heading(next(counter), 1)
    _quotation(document, data, labels, language)

    if data.get("wbs") or data.get("schedule"):
        appendix = document.add_section(WD_SECTION_START.NEW_PAGE)
        appendix.different_first_page_header_footer = False
        if data.get("wbs"):
            document.add_heading(next(counter), 1)
            _wbs_appendix(document, data, labels, language)
        if data.get("schedule"):
            document.add_heading(next(counter), 1)
            _schedule_appendix(document, data, labels, language)

    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()
