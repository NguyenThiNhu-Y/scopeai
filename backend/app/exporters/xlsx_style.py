"""Shared openpyxl styling so every exported workbook looks consistent."""

from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FILL = PatternFill(fill_type="solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
PARENT_FILL = PatternFill(fill_type="solid", fgColor="F2F2F2")
TOTAL_FILL = PatternFill(fill_type="solid", fgColor="D9E2F3")
ANSWER_FILL = PatternFill(fill_type="solid", fgColor="FFFF00")
TITLE_FONT = Font(bold=True, size=14, color="1F3A5F")
WRAP = Alignment(wrap_text=True, vertical="top")
PHASE_FILLS = {
    "poc": PatternFill(fill_type="solid", fgColor="8FAADC"),
    "mvp": PatternFill(fill_type="solid", fgColor="A9D18E"),
    "production": PatternFill(fill_type="solid", fgColor="F4B183"),
}


def money_format(currency: str | None) -> str:
    return "#,##0.00" if currency == "USD" else "#,##0"


def header_row(ws, row: int, headers: list[str], start_col: int = 1) -> None:
    for offset, title in enumerate(headers):
        cell = ws.cell(row, start_col + offset, title)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def body_cells(ws, first_row: int, last_row: int, first_col: int, last_col: int, wrap: bool = False) -> None:
    for row in ws.iter_rows(min_row=first_row, max_row=last_row, min_col=first_col, max_col=last_col):
        for cell in row:
            cell.border = BORDER
            if wrap:
                cell.alignment = WRAP


def fill_row(ws, row: int, first_col: int, last_col: int, fill: PatternFill, bold: bool = False) -> None:
    for col in range(first_col, last_col + 1):
        cell = ws.cell(row, col)
        cell.fill = fill
        if bold:
            cell.font = Font(bold=True)


def widths(ws, values: list[float]) -> None:
    for index, width in enumerate(values, 1):
        ws.column_dimensions[get_column_letter(index)].width = width


def title(ws, text: str, span: int) -> None:
    ws.cell(1, 1, text).font = TITLE_FONT
    if span > 1:
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=span)
