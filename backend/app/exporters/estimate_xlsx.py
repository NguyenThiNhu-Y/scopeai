from __future__ import annotations

from datetime import timedelta
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from app.exporters import xlsx_style as st
from app.exporters.common import (
    PHASES,
    PRICE_PLACEHOLDER,
    ROLE_LABELS,
    as_date,
    phase_label,
    role_label,
    wbs_rows,
    week_starts,
)

IMPACT_VI = {"solution": "Giải pháp", "effort": "Effort", "price": "Giá"}
STATUS_VI = {"awaiting_answer": "Chờ trả lời", "answered": "Đã trả lời", "skipped": "Bỏ qua"}
QA_HEADERS = ["Mã câu hỏi", "Vòng", "Câu hỏi", "Lý do (tiếng Việt)", "Ảnh hưởng", "Trả lời", "Là giả định", "Trạng thái"]
MD_FORMAT = "#,##0.0"
AMOUNT = "Thành tiền"


def _wbs_sheet(wb, data: dict) -> int:
    ws = wb.create_sheet("WBS")
    headers = ["Mã", "Công việc", "Giai đoạn", "Cấp", "Vai trò", "Man-day", "Requirement"]
    st.header_row(ws, 1, headers)
    row = 1
    for item, level, is_parent in wbs_rows(list(data.get("wbs") or [])):
        row += 1
        codes = getattr(item, "requirement_codes", None) or []
        ws.append([
            item.task_code,
            getattr(item, "name_vi", ""),
            phase_label(getattr(item, "phase", None), "vi"),
            level,
            None if is_parent else role_label(getattr(item, "role", None)),
            None if is_parent else float(item.man_day or 0),
            ", ".join(codes),
        ])
        ws.cell(row, 2).alignment = Alignment(indent=level - 1, wrap_text=True, vertical="top")
        ws.cell(row, 6).number_format = "#,##0.0"
        if is_parent:
            st.fill_row(ws, row, 1, len(headers), st.PARENT_FILL, bold=True)
    last = max(row, 2)
    st.body_cells(ws, 2, last, 1, len(headers))
    total_row = last + 1
    ws.cell(total_row, 2, "Tổng man-day")
    ws.cell(total_row, 6, f"=SUM(F2:F{last})").number_format = "#,##0.0"
    st.fill_row(ws, total_row, 1, len(headers), st.TOTAL_FILL, bold=True)
    for offset, phase in enumerate(PHASES, 1):
        label = phase_label(phase, "vi")
        ws.cell(total_row + offset, 2, f"Man-day {label}")
        ws.cell(total_row + offset, 3, label)
        ws.cell(total_row + offset, 6, f'=SUMIF(C2:C{last},"{label}",F2:F{last})').number_format = "#,##0.0"
    st.body_cells(ws, total_row, total_row + len(PHASES), 1, len(headers))
    st.widths(ws, [10, 48, 12, 6, 13, 10, 24])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:G{last}"
    return total_row


def _quote_lines(detail: dict) -> list[tuple[str, float | None, float | None, float]]:
    """Line items whose amounts sum exactly to the code-computed grand total."""
    lines = []
    by_role = detail.get("by_role") or {}
    for role in [r for r in ROLE_LABELS if r in by_role] + [r for r in by_role if r not in ROLE_LABELS]:
        value = by_role[role]
        lines.append((role_label(role), float(value.get("man_day") or 0), float(value.get("rate") or 0), float(value.get("amount") or 0)))
    overhead = detail.get("overhead") or {}
    extras = [
        ("Overhead PM", overhead.get("pm")),
        ("Overhead BrSE", overhead.get("brse")),
        ("Dự phòng rủi ro (contingency)", detail.get("contingency_amount")),
        ("Chi phí vận hành", detail.get("operating_cost")),
    ]
    lines += [(name, None, None, float(amount)) for name, amount in extras if amount]
    adjustment = float(detail.get("grand_total") or 0) - sum(line[3] for line in lines)
    if abs(adjustment) >= 0.005:
        lines.append(("Điều chỉnh giá (price-to-win / làm tròn)", None, None, round(adjustment, 2)))
    return lines


def _money(value, approved: bool):
    return value if approved else PRICE_PLACEHOLDER


def _quote_lines_table(ws, head: int, detail: dict, approved: bool, fmt: str) -> int:
    st.header_row(ws, head, ["Hạng mục", "Man-day", "Đơn giá", AMOUNT])
    lines = (_quote_lines(detail) if detail else []) or [("Giá trị trước duyệt", None, None, 0.0)]
    first = head + 1
    for offset, (name, man_day, rate, amount) in enumerate(lines):
        row = first + offset
        ws.cell(row, 1, name)
        ws.cell(row, 2, man_day).number_format = MD_FORMAT
        if rate is not None:
            ws.cell(row, 3, _money(rate, approved)).number_format = fmt
        ws.cell(row, 4, _money(amount, approved)).number_format = fmt
    last = first + len(lines) - 1
    total_row = last + 1
    ws.cell(total_row, 1, "Tổng (gồm thuế)" if detail.get("tax_mode") == "tax_included" else "Tổng (chưa thuế)")
    ws.cell(total_row, 2, f"=SUM(B{first}:B{last})").number_format = MD_FORMAT
    ws.cell(total_row, 4, _money(f"=SUM(D{first}:D{last})", approved)).number_format = fmt
    st.body_cells(ws, first, total_row, 1, 4)
    st.fill_row(ws, total_row, 1, 4, st.TOTAL_FILL, bold=True)
    return total_row


def _phase_table(ws, row: int, detail: dict, approved: bool, fmt: str) -> int:
    by_phase = detail.get("by_phase") or {}
    phase_days = detail.get("phase_man_days") or {}
    phases = [p for p in PHASES if p in by_phase or phase_days.get(p)]
    if not phases:
        return row
    st.header_row(ws, row, ["Giai đoạn", "Man-day", "", AMOUNT])
    for offset, phase in enumerate(phases, 1):
        ws.cell(row + offset, 1, phase_label(phase, "vi"))
        ws.cell(row + offset, 2, float(phase_days.get(phase) or 0)).number_format = MD_FORMAT
        ws.cell(row + offset, 4, _money(by_phase.get(phase), approved)).number_format = fmt
    st.body_cells(ws, row + 1, row + len(phases), 1, 4)
    return row + len(phases) + 2


def _payment_table(ws, row: int, data: dict, detail: dict, total_row: int, approved: bool, fmt: str) -> None:
    plan = detail.get("payment_plan") or []
    items = plan or (data.get("schedule") or {}).get("payment_milestones") or []
    if not items:
        return
    decimals = 2 if data.get("currency") == "USD" else 0
    st.header_row(ws, row, ["Mốc thanh toán", "Tỷ lệ (%)", "Ngày / tháng", AMOUNT])
    for offset, item in enumerate(items, 1):
        current = row + offset
        ws.cell(current, 1, item.get("name", ""))
        ws.cell(current, 2, float(item.get("percent") or 0)).number_format = "0.00"
        ws.cell(current, 3, item.get("due_date") or (f"Tháng {item['month']}" if item.get("month") else ""))
        # Fixed-price milestones only carry a percent; Excel derives the amount from the total cell.
        amount = item.get("amount") if plan else f"=ROUND($D${total_row}*B{current}/100,{decimals})"
        ws.cell(current, 4, _money(amount, approved)).number_format = fmt
    st.body_cells(ws, row + 1, row + len(items), 1, 4)


def _quotation_sheet(wb, data: dict) -> int | None:
    ws = wb.create_sheet("Báo giá")
    approved = bool(data.get("approved"))
    detail = data.get("quotation_detail") or {}
    currency = data.get("currency") or ""
    fmt = st.money_format(currency)
    st.title(ws, f"Báo giá — {data['project_name']}", 4)
    info = [
        ("Hình thức hợp đồng", detail.get("contract_model") or ""),
        ("Tiền tệ", currency),
        ("Thuế", "Đã gồm thuế" if detail.get("tax_mode") == "tax_included" else "Chưa gồm thuế"),
        ("Hiệu lực đến", str(detail.get("valid_until") or "")),
        ("Cổng giá", "Đã duyệt" if approved else "Chưa duyệt — chưa hiển thị giá"),
    ]
    for offset, (key, value) in enumerate(info, 3):
        ws.cell(offset, 1, key).font = Font(bold=True)
        ws.cell(offset, 2, value)
    total_row = _quote_lines_table(ws, 3 + len(info) + 1, detail, approved, fmt)
    next_row = _phase_table(ws, total_row + 2, detail, approved, fmt)
    _payment_table(ws, next_row, data, detail, total_row, approved, fmt)
    st.widths(ws, [40, 12, 18, 22])
    return total_row if approved else None


def _gantt_sheet(wb, data: dict) -> None:
    ws = wb.create_sheet("Lịch Gantt")
    schedule = data.get("schedule") or {}
    tasks = [t for t in schedule.get("tasks") or [] if as_date(t.get("start_date")) and as_date(t.get("end_date"))]
    names = {item.task_code: getattr(item, "name_vi", "") for item in data.get("wbs") or []}
    fixed = ["Mã", "Công việc", "Giai đoạn", "Vai trò", "Bắt đầu", "Kết thúc"]
    milestones = [(f"◆ {m.get('name', '')} ({float(m.get('percent') or 0):g}%)", as_date(m.get("due_date"))) for m in schedule.get("payment_milestones") or []]
    if data.get("customer_deadline"):
        milestones.append(("▲ Hạn khách yêu cầu", as_date(data["customer_deadline"])))
    milestones = [(name, day) for name, day in milestones if day]
    dates = [as_date(t["start_date"]) for t in tasks] + [as_date(t["end_date"]) for t in tasks] + [day for _, day in milestones]
    weeks = week_starts(min(dates), max(dates)) if dates else []
    st.header_row(ws, 1, fixed + [f"W{index:02d}\n{day:%d/%m}" for index, day in enumerate(weeks, 1)])
    ws.row_dimensions[1].height = 30
    first_week_col = len(fixed) + 1

    def week_col(day) -> int | None:
        for index, start in enumerate(weeks):
            if start <= day < start + timedelta(days=7):
                return first_week_col + index
        return None

    row = 1
    for task in sorted(tasks, key=lambda t: (as_date(t["start_date"]), t.get("task_code", ""))):
        row += 1
        start, end = as_date(task["start_date"]), as_date(task["end_date"])
        ws.append([task.get("task_code", ""), names.get(task.get("task_code"), ""), phase_label(task.get("phase"), "vi"), role_label(task.get("role")), start, end])
        ws.cell(row, 5).number_format = ws.cell(row, 6).number_format = "yyyy-mm-dd"
        fill = st.PHASE_FILLS.get(task.get("phase"))
        for index, week in enumerate(weeks):
            if week <= end and start < week + timedelta(days=7) and fill is not None:
                ws.cell(row, first_week_col + index).fill = fill
    if milestones:
        row += 1
    for name, day in milestones:
        row += 1
        ws.cell(row, 2, name).font = Font(bold=True)
        ws.cell(row, 5, day).number_format = "yyyy-mm-dd"
        col = week_col(day)
        if col:
            marker = ws.cell(row, col, name[0])
            marker.font = Font(bold=True, color="C00000")
            marker.alignment = Alignment(horizontal="center")
    if not tasks:
        ws.cell(2, 2, "Chưa có lịch — hãy tính lịch ở bước Schedule trước khi xuất.")
    st.body_cells(ws, 2, max(row, 2), 1, len(fixed))
    st.widths(ws, [10, 40, 12, 12, 12, 12])
    for index in range(len(weeks)):
        ws.column_dimensions[get_column_letter(first_week_col + index)].width = 6.5
    ws.freeze_panes = ws.cell(2, first_week_col)


def write_internal_qa(ws, questions) -> None:
    st.header_row(ws, 1, QA_HEADERS)
    row = 1
    for q in questions:
        row += 1
        ws.append([
            q.id, q.round_no, q.question, q.reason_vi, IMPACT_VI.get(q.impact, q.impact),
            q.answer or "", "Có" if q.is_assumption else "Không", STATUS_VI.get(q.status, q.status),
        ])
        ws.cell(row, 6).fill = st.ANSWER_FILL
    st.body_cells(ws, 2, max(row, 2), 1, len(QA_HEADERS), wrap=True)
    st.widths(ws, [11, 7, 50, 40, 12, 40, 11, 13])
    ws.freeze_panes = "A2"


def _summary_sheet(ws, data: dict, wbs_total_row: int, quote_total_row: int | None) -> None:
    ws.title = "Tóm tắt"
    schedule = data.get("schedule") or {}
    st.title(ws, f"Bảng ước tính — {data['project_name']}", 2)
    rows = [
        ("Dự án", data["project_name"]),
        ("Khách hàng", data.get("customer_name", "")),
        ("Ngôn ngữ khách", data.get("customer_language", "")),
        ("Tiền tệ", data.get("currency", "")),
        ("Ngày xuất", data.get("export_date", "")),
        ("Cổng giá", "Đã duyệt" if data.get("approved") else "Chưa duyệt"),
        ("Tổng man-day", f"=WBS!F{wbs_total_row}"),
        ("Tổng báo giá", f"='Báo giá'!D{quote_total_row}" if quote_total_row else PRICE_PLACEHOLDER),
        ("Ngày bắt đầu", str(schedule.get("start_date") or "")),
        ("Ngày kết thúc", str(schedule.get("end_date") or "")),
    ]
    for offset, (key, value) in enumerate(rows, 3):
        ws.cell(offset, 1, key).font = Font(bold=True)
        ws.cell(offset, 2, value)
    ws.cell(9, 2).number_format = "#,##0.0"
    ws.cell(10, 2).number_format = st.money_format(data.get("currency"))
    st.body_cells(ws, 3, 2 + len(rows), 1, 2)
    st.widths(ws, [22, 50])


def render_estimate(data: dict) -> bytes:
    wb = Workbook()
    summary = wb.active
    assert summary is not None
    wbs_total_row = _wbs_sheet(wb, data)
    quote_total_row = _quotation_sheet(wb, data)
    _gantt_sheet(wb, data)
    write_internal_qa(wb.create_sheet("Q&A"), data.get("questions") or [])
    _summary_sheet(summary, data, wbs_total_row, quote_total_row)
    stream = BytesIO()
    wb.save(stream)
    return stream.getvalue()
