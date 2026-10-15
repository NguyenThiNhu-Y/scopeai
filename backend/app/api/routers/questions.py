"""Router hỏi làm rõ — contract api_contract.md §6 (M2)."""

from __future__ import annotations

from datetime import date, datetime, timezone
from zipfile import BadZipFile
from fastapi import APIRouter, Depends, Query, UploadFile, File
from fastapi.responses import Response
from io import BytesIO
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Protection
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy.orm import Session

from app.api.deps import Identity, get_db, get_identity
from app.api.errors import state_conflict
from app.exporters import xlsx_style as st
from app.exporters.filename import content_disposition, export_filename
from app.models import ExportLog
from app.pipeline import is_running, resume
from app.schemas.api_schemas.question import (
    AnswerItemIn,
    AnswersIn,
    AnswersOut,
    QuestionOut,
    QuestionsOut,
)
from app.services.clarify_service import list_questions, submit_answers
from app.services.bid_service import get_bid_or_404
from app.llm.gateway import complete_with_retry, get_llm_client
from app.schemas.llm_schemas.email_draft import EmailDraft
from app.privacy import mask_text
from pydantic import BaseModel

MAX_QA_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_QA_ROWS = 200

router = APIRouter(prefix="/bids", tags=["questions"])


def _customer_reason(language: str, impact: str, fallback: str) -> str:
    """Use stable customer-facing labels without asking the LLM to translate."""
    labels = {
        "ja": {"effort": "作業量に影響するため", "solution": "提案内容に影響するため", "price": "費用に影響するため"},
        "en": {"effort": "Because it affects the effort", "solution": "Because it affects the proposed solution", "price": "Because it affects the cost"},
        "vi": {"effort": "Vì ảnh hưởng đến effort", "solution": "Vì ảnh hưởng đến phương án", "price": "Vì ảnh hưởng đến chi phí"},
    }
    return labels.get(language, labels["vi"]).get(impact, fallback)

QA_TEXT = {
    "ja": {
        "headers": ["質問コード", "ラウンド", "質問", "理由", "回答"],
        "guide_sheet": "記入方法", "title": "ご質問事項（Q&A）", "project": "案件名", "customer": "お客様", "date": "発行日",
        "steps": ["黄色の「回答」列のみご記入ください。", "質問コード・ラウンドは変更しないでください（取込時の照合に使用します）。", "未確定の項目は空欄のままで構いません。"],
    },
    "en": {
        "headers": ["Question code", "Round", "Question", "Reason", "Answer"],
        "guide_sheet": "Instructions", "title": "Clarification questions (Q&A)", "project": "Project", "customer": "Customer", "date": "Issued",
        "steps": ["Please fill in the yellow Answer column only.", "Do not change Question code or Round; they are used to match your answers.", "Leave the answer blank if it is not decided yet."],
    },
    "vi": {
        "headers": ["Mã câu hỏi", "Vòng", "Câu hỏi", "Lý do", "Trả lời"],
        "guide_sheet": "Hướng dẫn", "title": "Câu hỏi làm rõ (Q&A)", "project": "Dự án", "customer": "Khách hàng", "date": "Ngày phát hành",
        "steps": ["Vui lòng chỉ điền cột Trả lời (màu vàng).", "Không sửa Mã câu hỏi và Vòng — dùng để đối chiếu khi nhập lại.", "Để trống nếu chưa có câu trả lời."],
    },
}


@router.get("/{bid_id}/questions/export")
def export_questions(bid_id: int, session: Session = Depends(get_db), identity: Identity = Depends(get_identity)):
    bid = get_bid_or_404(session, bid_id)
    rows = list_questions(session, bid_id)["items"]
    lang = bid.customer_language if bid.customer_language in QA_TEXT else "vi"
    text = QA_TEXT[lang]
    today = date.today().isoformat()
    wb = Workbook(); ws = wb.active; ws.title = "Q&A"
    st.header_row(ws, 1, text["headers"])
    for q in rows:
        ws.append([q.id, q.round_no, q.question, _customer_reason(lang, q.impact, q.reason_vi), q.answer or ""])
        answer = ws.cell(ws.max_row, 5)
        answer.fill = st.ANSWER_FILL
        answer.protection = Protection(locked=False)
    st.body_cells(ws, 2, max(ws.max_row, 2), 1, 5, wrap=True)
    st.widths(ws, [14, 9, 60, 30, 50])
    ws.freeze_panes = "A2"
    # Lock everything but the answer column so codes/rounds cannot be edited by mistake (no password).
    ws.protection.sheet = True
    ws.protection.formatColumns = False
    ws.protection.formatRows = False

    guide = wb.create_sheet(text["guide_sheet"])
    st.title(guide, text["title"], 2)
    for offset, (key, value) in enumerate([(text["project"], bid.project_name), (text["customer"], bid.customer_name), (text["date"], today)], 3):
        guide.cell(offset, 1, key).font = Font(bold=True)
        guide.cell(offset, 2, value)
    for offset, step in enumerate(text["steps"], 7):
        guide.cell(offset, 1, f"{offset - 6}.")
        guide.cell(offset, 2, step)
    st.widths(guide, [16, 90])

    out = BytesIO(); wb.save(out)
    session.add(ExportLog(bid_id=bid.id, kind="qa_xlsx", actor=identity.user, price_gate="n/a", ja_check=None)); session.commit()
    filename = export_filename(bid.customer_name, bid.project_name, "QA", export_date=today, ext="xlsx")
    return Response(out.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": content_disposition(filename)})

@router.post("/{bid_id}/questions/import", response_model=AnswersOut)
async def import_questions(bid_id: int, file: UploadFile = File(...), session: Session = Depends(get_db), identity: Identity = Depends(get_identity)):
    get_bid_or_404(session, bid_id)
    try:
        payload = await file.read(MAX_QA_UPLOAD_BYTES + 1)
        if len(payload) > MAX_QA_UPLOAD_BYTES:
            from app.api.errors import validation_error
            raise validation_error("Q&A sheet vượt quá giới hạn 5 MB", {"max_bytes": MAX_QA_UPLOAD_BYTES})
        book = load_workbook(BytesIO(payload), read_only=True, data_only=True)
        # The customer may save the file while viewing the instructions sheet.
        ws = book["Q&A"] if "Q&A" in book.sheetnames else book.active
        headers = [str(value or "").strip().casefold() for value in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
        aliases = {
            "code": {"mã câu hỏi", "question code", "質問コード"},
            "round": {"vòng", "round", "ラウンド", "質問ラウンド"},
            "answer": {"trả lời", "answer", "回答"},
        }
        column = {}
        for key, candidates in aliases.items():
            match = next((index for index, header in enumerate(headers) if header in candidates), None)
            if match is not None:
                column[key] = match
        if len(column) != len(aliases):
            from app.api.errors import validation_error
            raise validation_error("Q&A sheet thiếu cột bắt buộc: Mã câu hỏi, Vòng, Trả lời")
        rows = []
        for row in ws.iter_rows(min_row=2, max_row=MAX_QA_ROWS + 2, values_only=True):
            rows.append(row)
        if len(rows) > MAX_QA_ROWS:
            from app.api.errors import validation_error
            raise validation_error("Q&A sheet vượt quá giới hạn 200 dòng", {"max_rows": MAX_QA_ROWS})
    except (BadZipFile, InvalidFileException, TypeError, ValueError, IndexError, KeyError) as exc:
        from app.api.errors import validation_error
        raise validation_error("Q&A sheet không hợp lệ", {"reason": str(exc)})
    from app.models import Question
    questions = {question.id: question for question in session.query(Question).filter_by(bid_id=bid_id).all()}
    answers_by_round = {}
    warnings = []
    for row_number, row in enumerate(rows, start=2):
        code = row[column["code"]] if len(row) > column["code"] else None
        round_value = row[column["round"]] if len(row) > column["round"] else None
        answer_value = row[column["answer"]] if len(row) > column["answer"] else None
        if not str(answer_value or "").strip():
            warnings.append(f"Dòng {row_number}: bỏ qua ô trả lời trống")
            continue
        try:
            question_id = int(code)
            round_no = int(round_value)
        except (TypeError, ValueError):
            warnings.append(f"Dòng {row_number}: mã câu hỏi hoặc vòng không hợp lệ")
            continue
        question = questions.get(question_id)
        if question is None:
            warnings.append(f"Dòng {row_number}: mã câu hỏi không tồn tại")
            continue
        if question.round_no != round_no:
            warnings.append(f"Dòng {row_number}: mã câu hỏi không thuộc vòng đã ghi")
            continue
        if question.status != "awaiting_answer":
            warnings.append(f"Dòng {row_number}: câu hỏi vòng {round_no} đã đóng, bỏ qua")
            continue
        answers_by_round.setdefault(round_no, []).append({"question_id": question_id, "answer": str(answer_value).strip(), "is_assumption": False})
    if not answers_by_round:
        bid = get_bid_or_404(session, bid_id)
        return AnswersOut(status=bid.status, pipeline_resumed_from="gap", warnings=warnings)
    active_rounds = {
        question.round_no
        for question in questions.values()
        if question.status == "awaiting_answer"
    }
    result = None
    candidate_rounds = active_rounds & answers_by_round.keys()
    if candidate_rounds:
        round_no = max(candidate_rounds)
        for older_round in sorted(candidate_rounds):
            if older_round != round_no:
                warnings.append(f"Vòng {older_round}: chỉ xử lý vòng đang chờ mới nhất")
        awaiting = {
            question.id
            for question in questions.values()
            if question.round_no == round_no and question.status == "awaiting_answer"
        }
        submitted = {answer["question_id"] for answer in answers_by_round[round_no]}
        for question_id in sorted(awaiting - submitted):
            warnings.append(f"Câu hỏi {question_id} vòng {round_no}: chưa có câu trả lời")
        if submitted.issuperset(awaiting):
            result = submit_answers(session, bid_id, round_no, answers_by_round[round_no], answered_by=identity.user)
        else:
            for answer in answers_by_round[round_no]:
                question = questions[answer["question_id"]]
                question.answer = answer["answer"]
                question.answered_by = identity.user
                question.answered_at = datetime.now(timezone.utc)
                question.status = "answered"
            session.commit()
    if result is None:
        bid = get_bid_or_404(session, bid_id)
        return AnswersOut(status=bid.status, pipeline_resumed_from="gap", warnings=warnings)
    resume(session, bid_id, from_step="gap")
    return AnswersOut(**result, warnings=warnings)


class EmailDraftOut(BaseModel):
    """Editable customer email with identity fields owned by the application."""

    bid_id: int
    customer_name: str
    sender: str
    customer_language: str
    subject: str
    body: str


@router.post("/{bid_id}/questions/email-draft", response_model=EmailDraftOut)
def email_draft(
    bid_id: int,
    session: Session = Depends(get_db),
    identity: Identity = Depends(get_identity),
) -> EmailDraftOut:
    """Sinh email nháp theo ngôn ngữ khách; tên khách/người gửi do code điền."""
    bid = get_bid_or_404(session, bid_id)
    prompt = mask_text(
        f"Write a concise clarification email in {bid.customer_language}. "
        f"Customer name: {bid.customer_name}. Sender: {identity.user}."
    )
    result = complete_with_retry(
        get_llm_client(), EmailDraft, prompt,
        context={"customer_language": bid.customer_language},
    )
    return EmailDraftOut(
        bid_id=bid.id,
        customer_name=bid.customer_name,
        sender=identity.user,
        customer_language=bid.customer_language,
        subject=result["subject"].replace("{{CUSTOMER_NAME}}", bid.customer_name),
        body=result["body"].replace("{{CUSTOMER_NAME}}", bid.customer_name),
    )


def _to_question_out(q) -> QuestionOut:
    return QuestionOut(
        id=q.id,
        round_no=q.round_no,
        question=q.question,
        reason_vi=q.reason_vi,
        impact=q.impact,
        related_req_codes=q.related_req_codes or [],
        answer=q.answer,
        is_assumption=q.is_assumption,
        status=q.status,
    )


@router.get("/{bid_id}/questions", response_model=QuestionsOut)
def get_questions(
    bid_id: int,
    round: int | None = Query(default=None, ge=1, le=2),
    session: Session = Depends(get_db),
    identity: Identity = Depends(get_identity),
) -> QuestionsOut:
    """Danh sách câu hỏi làm rõ (FR-007.2, FR-010, §6.1)."""
    data = list_questions(session, bid_id, round_no=round)
    return QuestionsOut(
        items=[_to_question_out(q) for q in data["items"]],
        rounds_used=data["rounds_used"],
        max_rounds=data["max_rounds"],
        round_two_reason_vi=data["round_two_reason_vi"],
    )


@router.post("/{bid_id}/questions/{round}/answers", response_model=AnswersOut)
async def post_answers(
    bid_id: int,
    round: int,
    body: AnswersIn,
    session: Session = Depends(get_db),
    identity: Identity = Depends(get_identity),
) -> AnswersOut:
    """Trả lời câu hỏi (FR-011, §6.2) — thiếu câu → 400 kèm details.missing_ids.

    Trả lời đủ → pipeline tự chạy tiếp từ bước ``gap`` (không làm lại intake).
    """
    # Chặn trước khi lưu answers: nếu pipeline vẫn đang chạy (race hiếm) →
    # 409 STATE_CONFLICT, tránh lưu answers rồi resume thất bại (M3 review).
    if is_running(bid_id):
        raise state_conflict("Pipeline đang chạy — vui lòng chờ trước khi trả lời")

    answers = [a.model_dump() for a in body.answers]
    result = submit_answers(session, bid_id, round, answers, answered_by=identity.user)

    # Pipeline tự chạy tiếp từ gap — bước intake đã completed được giữ nguyên.
    resume(session, bid_id, from_step="gap")

    return AnswersOut(**result)
